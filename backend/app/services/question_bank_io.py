"""题库批量导入导出：老师自助录题（2026-10-14 本期范围）。

数据流：模板下载 → 老师填 Excel/CSV → import-preview（解析+校验+查重，
不落库）→ import-confirm（重新上传同一文件，单事务落库）。数据只写现有
四张题型表（Passage / RepeatSentence / ScenarioQuestion / Instruction），
零迁移、零新表。

口径：
- 只 INSERT 不 UPDATE：导入永不覆盖旧题（发布快照与历史成绩不受影响的
  根基），改题仍走现有编辑页；
- 校验与手动建题同一套：validate_exam_fields /
  validate_question_suggested_seconds / band 白名单 / 字段长度与秒数区间；
- 重复两级（规范化 = NFKC + strip + casefold + 空白折叠）：
  完全重复（题干与配置都相同，含批内）默认跳过；疑似重复（题干相同、
  配置/主题不同）预览逐行列出，确认时由 include_suspected 决定；
- 兜底：问答题主题不存在 → 自动新建 Scenario（预览明示）；篇目单元
  不匹配 → 降级为无单元；复述句挂靠篇目 0 命中 → 独立句、多命中 → 报错；
- 音频两步走：导入只收音频链接（须为本站 /api/v1/audio/content/ 前缀），
  音频文件上传沿用现有入口；缺音频按题型计数提示，不阻断——「题目已
  导入」不等于「音频已齐备」。

xlsx 模板与解析共用本文件的列规格（单一事实源）；CSV 与导出格式同构，
导出→导入往返不丢教师可填的关键配置（拆句/逐题计时不入模板，属声明的
例外，模板说明页注明）。
"""

import csv
import io
import unicodedata
import uuid
from dataclasses import dataclass, field
from typing import Any

from fastapi import HTTPException
from sqlmodel import Session, col, select

from app import crud
from app.core.storage import CONTENT_AUDIO_URL_PREFIX
from app.models import (
    EXAM_KIND_QUESTION,
    EXAM_KIND_REPEAT,
    VOCAB_LEVEL_ORDER,
    Instruction,
    Passage,
    QuestionBankImportIssue,
    QuestionBankImportPreview,
    QuestionBankImportResult,
    QuestionBankImportTypeResult,
    QuestionBankTypeCount,
    RepeatSentence,
    Scenario,
    ScenarioQuestion,
    Unit,
    get_datetime_utc,
    validate_exam_fields,
    validate_question_suggested_seconds,
)

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
# 导出上限必须与导入上限一致：导出的文件要始终能直接回导（评审 P2 曾出现
# 2001 行可导出、回导被 2000 行上限拒绝）。更大规模请按题型/主题筛选分批导出
MAX_EXPORT_ROWS = MAX_IMPORT_ROWS
MAX_ISSUES = 200
MAX_SUSPECTED_LISTED = 100
# csv 模块单字段默认 128KB 上限：长正文导出 CSV 后回导会触发 csv.Error → 500。
# 抬到上传上限之上，任何能通过 5MB 限制的文件都可解析（_parse_csv 另有兜底 422）
csv.field_size_limit(MAX_UPLOAD_BYTES + 1024)
# xlsx 解析资源上限（评审 P1）：按工作表声明范围在读遍历前判定，
# 防止「只有格式没有内容」的远端单元格把遍历量撑大（表头+说明行余量）
MAX_SHEET_ROWS = MAX_IMPORT_ROWS + 8
MAX_SHEET_COLUMNS = 64
# Excel 单元格字符串硬上限：openpyxl 写超长会静默截断，导出前显式拒绝
EXCEL_CELL_MAX_CHARS = 32767

TYPE_PASSAGE = "passage"
TYPE_REPEAT = "repeat"
TYPE_QUESTION = "question"
TYPE_INSTRUCTION = "instruction"
QUESTION_BANK_TYPE_ORDER = (TYPE_PASSAGE, TYPE_REPEAT, TYPE_QUESTION, TYPE_INSTRUCTION)

TYPE_LABELS = {
    TYPE_PASSAGE: "文章朗读",
    TYPE_REPEAT: "听句复述",
    TYPE_QUESTION: "情景问答",
    TYPE_INSTRUCTION: "题目说明",
}

# 工作表名 → 题型（模板工作表名/CSV 题型列共用；大小写与首尾空白不敏感）
SHEET_NAME_ALIASES: dict[str, set[str]] = {
    TYPE_PASSAGE: {
        "文章朗读",
        "文章朗读 · 篇目",
        "朗读",
        "read aloud",
        "passage",
        "reading",
    },
    TYPE_REPEAT: {
        "听句复述",
        "听令复述",
        "listen & repeat",
        "listen and repeat",
        "repeat",
    },
    TYPE_QUESTION: {
        "情景问答",
        "问答",
        "scenario q&a",
        "scenario qa",
        "question",
        "qa",
    },
    # 注意：「说明」不进题型别名——那是模板说明页的名字，不能当题目说明工作表
    TYPE_INSTRUCTION: {"题目说明", "instruction"},
}
# 说明页：有内容也整体跳过，不参与导入（「说明」从题型别名挪到这里）
INSTRUCTION_SHEET_ALIASES = {
    "说明",
    "填写说明",
    "使用说明",
    "说明页",
    "readme",
    "instructions",
    "help",
}

VALID_BANDS = {"A2", "B1", "B2"}

# 考试题型别名（值匹配用：casefold + 空白折叠后比对；含系统代码/中英文名）
EXAM_KIND_ALIASES: dict[str, set[str]] = {
    "toefl_lnr": {
        "toefl_lnr",
        "toefl",
        "托福",
        "toefl · 听令复述",
        "toefl听令复述",
        "toefl · 听句复述",
        "toefl听句复述",
        "listen & repeat",
        "listen and repeat",
    },
    "interview": {
        "interview",
        "toefl · 情景访谈",
        "toefl情景访谈",
        "take an interview",
        "访谈",
    },
    "ielts_p1": {"ielts_p1", "ielts · part 1", "ielts part 1", "ielts p1", "part 1"},
    "ielts_p2": {"ielts_p2", "ielts · part 2", "ielts part 2", "ielts p2", "part 2"},
    "ielts_p3": {"ielts_p3", "ielts · part 3", "ielts part 3", "ielts p3", "part 3"},
}
EXAM_LEVEL_ALIASES: dict[str, set[str]] = {
    "KET": {"ket", "ket 词汇"},
    "PET": {"pet", "pet 词汇"},
    "ACADEMIC": {"academic", "学术词汇", "学术"},
    "CET4": {"cet4", "cet-4", "cet 4", "四级"},
    "IELTS_TOEFL": {
        "ielts_toefl",
        "雅思&托福词汇",
        "雅思托福",
        "雅思和托福",
        "ielts & toefl",
        "ielts&toefl",
    },
}
STATUS_ACTIVE_ALIASES = {"启用", "是", "true", "active", "1"}
STATUS_INACTIVE_ALIASES = {"停用", "否", "false", "inactive", "0"}

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv"


# ── 规范化小工具（表头匹配/查重键/单元格转字符串） ──────────────────


def _norm_header(raw: Any) -> str:
    """表头/枚举值规范化：casefold + 空白折叠（双语合并表头的换行也折叠）。"""
    if raw is None:
        return ""
    return " ".join(str(raw).split()).casefold()


def _norm_text(raw: str) -> str:
    """查重键规范化：NFKC（全角→半角）+ strip + casefold + 内部空白折叠。"""
    return " ".join(unicodedata.normalize("NFKC", raw).strip().casefold().split())


def _cell_str(raw: Any) -> str:
    """单元格值 → 字符串（openpyxl 可能给 int/float/bool/None）。"""
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return "是" if raw else "否"
    if isinstance(raw, float):
        return str(int(raw)) if raw.is_integer() else str(raw)
    return str(raw).strip()


# ── 列规格（canonical 字段 → 双语表头；模板生成与解析共用） ─────────


@dataclass(frozen=True)
class ColumnSpec:
    """一列的规格：模板表头单元格为「中文\n英文」，解析时三种形式都认。"""

    field: str
    zh: str
    en: str
    extra_aliases: tuple[str, ...] = ()

    @property
    def header(self) -> str:
        return f"{self.zh}\n{self.en}"

    def aliases(self) -> set[str]:
        return {
            _norm_header(a)
            for a in (self.zh, self.en, self.header, *self.extra_aliases)
        }


PASSAGE_COLUMNS = (
    ColumnSpec("title", "标题", "Title"),
    ColumnSpec("unit_title", "单元", "Unit"),
    ColumnSpec("topic", "主题", "Topic"),
    ColumnSpec("cefr_band", "难度", "Level", ("cefr", "cefr_band", "cefr level")),
    ColumnSpec("text", "正文", "Text", ("文章正文", "内容", "body")),
    ColumnSpec("translation", "中文提示", "Chinese hint", ("翻译", "hint")),
    ColumnSpec("suggested_seconds", "建议秒数", "Seconds"),
    ColumnSpec("audio_url", "音频链接", "Audio URL", ("音频", "audio")),
    ColumnSpec("is_active", "状态", "Status", ("启用", "active", "is_active")),
)
REPEAT_COLUMNS = (
    ColumnSpec("text", "英文句子", "Sentence", ("句子", "sentence")),
    ColumnSpec("translation", "中文提示", "Chinese hint", ("翻译", "hint")),
    ColumnSpec(
        "passage_title",
        "挂靠篇目",
        "Passage",
        ("所属篇目", "篇目标题", "passage title"),
    ),
    ColumnSpec("suggested_seconds", "作答秒数", "Seconds"),
    ColumnSpec(
        "replay_limit", "可听次数", "Replays", ("重听次数", "replay", "replays")
    ),
    ColumnSpec("exam_kind", "考试题型", "Exam type", ("exam_kind",)),
    ColumnSpec("exam_level", "考试级别", "Exam level", ("exam_level",)),
    ColumnSpec("audio_url", "音频链接", "Audio URL", ("音频", "audio")),
)
QUESTION_COLUMNS = (
    ColumnSpec("topic", "主题", "Topic"),
    ColumnSpec("text", "英文题目", "Question", ("题目", "question")),
    ColumnSpec("translation", "中文提示", "Chinese hint", ("翻译", "hint")),
    ColumnSpec("suggested_seconds", "作答秒数", "Seconds"),
    ColumnSpec("exam_kind", "考试题型", "Exam type", ("exam_kind",)),
    ColumnSpec("exam_level", "考试级别", "Exam level", ("exam_level",)),
    ColumnSpec(
        "cue_card_bullets", "话题卡要点", "Cue card", ("话题卡", "cue_card_bullets")
    ),
    ColumnSpec(
        "prep_seconds", "准备秒数", "Prep seconds", ("准备时间", "prep", "prep_seconds")
    ),
    ColumnSpec("audio_url", "音频链接", "Audio URL", ("音频", "audio")),
)
INSTRUCTION_COLUMNS = (
    ColumnSpec("title", "小标题", "Title", ("标题",)),
    ColumnSpec("text", "说明文字", "Text", ("内容",)),
    ColumnSpec("suggested_seconds", "秒数", "Seconds"),
)

# CSV 单文件：题型列 + 四表列并集（列名跨题型复用同一 canonical 字段）
CSV_TYPE_COLUMN = ColumnSpec("__type__", "题型", "Type", ("item type", "kind"))
CSV_UNION_COLUMNS: tuple[ColumnSpec, ...] = (
    ColumnSpec("title", "标题", "Title", ("小标题",)),
    ColumnSpec("unit_title", "单元", "Unit"),
    ColumnSpec("topic", "主题", "Topic"),
    ColumnSpec("cefr_band", "难度", "Level", ("cefr",)),
    ColumnSpec(
        "text",
        "内容",
        "Text",
        (
            "正文",
            "英文句子",
            "句子",
            "英文题目",
            "题目",
            "说明文字",
            "sentence",
            "question",
            "body",
        ),
    ),
    ColumnSpec("translation", "中文提示", "Chinese hint", ("翻译", "hint")),
    ColumnSpec("passage_title", "挂靠篇目", "Passage", ("所属篇目", "篇目标题")),
    ColumnSpec(
        "suggested_seconds", "秒数", "Seconds", ("建议秒数", "作答秒数", "duration")
    ),
    ColumnSpec("replay_limit", "可听次数", "Replays", ("重听次数", "replay")),
    ColumnSpec("exam_kind", "考试题型", "Exam type", ("exam_kind",)),
    ColumnSpec("exam_level", "考试级别", "Exam level", ("exam_level",)),
    ColumnSpec("cue_card_bullets", "话题卡要点", "Cue card", ("话题卡",)),
    ColumnSpec("prep_seconds", "准备秒数", "Prep seconds", ("准备时间",)),
    ColumnSpec("audio_url", "音频链接", "Audio URL", ("音频", "audio")),
    ColumnSpec("is_active", "状态", "Status", ("启用", "active")),
)

COLUMNS_BY_TYPE: dict[str, tuple[ColumnSpec, ...]] = {
    TYPE_PASSAGE: PASSAGE_COLUMNS,
    TYPE_REPEAT: REPEAT_COLUMNS,
    TYPE_QUESTION: QUESTION_COLUMNS,
    TYPE_INSTRUCTION: INSTRUCTION_COLUMNS,
}
# 各题型必填列（canonical 字段）
REQUIRED_FIELDS: dict[str, tuple[str, ...]] = {
    TYPE_PASSAGE: ("title", "text"),
    TYPE_REPEAT: ("text",),
    TYPE_QUESTION: ("topic", "text"),
    TYPE_INSTRUCTION: ("text",),
}


def _build_header_lookup() -> dict[str, str]:
    """别名（规范化）→ canonical 字段。跨题型同名列（秒数/中文提示）指向同一字段。"""
    lookup: dict[str, str] = {}
    all_specs: list[ColumnSpec] = [CSV_TYPE_COLUMN, *CSV_UNION_COLUMNS]
    for specs in COLUMNS_BY_TYPE.values():
        all_specs.extend(specs)
    for spec in all_specs:
        for alias in spec.aliases():
            lookup.setdefault(alias, spec.field)
    return lookup


HEADER_LOOKUP = _build_header_lookup()


# ── 行级解析工具 ─────────────────────────────────────────────────────


class _RowError(Exception):
    """行级校验失败：reason 为中文稳定标识（前端做双语映射）。"""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


def _checked(call: Any) -> str | None:
    """跑共享校验函数（抛 HTTPException 422），把 detail 转成行级原因。"""
    try:
        call()
        return None
    except HTTPException as exc:
        return str(exc.detail)


def _parse_int(raw: str, *, default: int, lo: int, hi: int, label: str) -> int:
    s = (raw or "").strip()
    if not s:
        return default
    n: int | float
    try:
        n = int(s)
    except ValueError:
        try:
            f = float(s)
            n = int(f) if f.is_integer() else f
        except ValueError as exc:
            raise _RowError(f"{label}需为 {lo}–{hi} 之间的整数") from exc
    if not isinstance(n, int) or not lo <= n <= hi:
        raise _RowError(f"{label}需为 {lo}–{hi} 之间的整数")
    return n


def _parse_optional_text(raw: str, *, max_len: int, label: str) -> str | None:
    s = (raw or "").strip()
    if not s:
        return None
    if len(s) > max_len:
        raise _RowError(f"{label}过长（>{max_len} 字符）")
    return s


def _parse_band(raw: str) -> str:
    s = (raw or "").strip().upper()
    if not s:
        return "B1"
    if s not in VALID_BANDS:
        raise _RowError("难度无效，可选：A2/B1/B2")
    return s


def _parse_status(raw: str) -> bool:
    s = _norm_header(raw)
    if not s:
        return True
    if s in STATUS_ACTIVE_ALIASES:
        return True
    if s in STATUS_INACTIVE_ALIASES:
        return False
    raise _RowError("状态无效，可选：启用/停用")


def _parse_audio(raw: str) -> str | None:
    s = (raw or "").strip()
    if not s:
        return None
    path = s
    for scheme in ("https://", "http://"):
        if path.startswith(scheme):
            rest = path[len(scheme) :]
            path = rest.split("/", 1)[1] if "/" in rest else ""
            break
    if not path.startswith(f"{CONTENT_AUDIO_URL_PREFIX}/"):
        raise _RowError(
            f"音频链接格式不正确（应为本站 {CONTENT_AUDIO_URL_PREFIX}/ 开头的链接，"
            "音频文件请在导入后到题目编辑里上传）"
        )
    return s


def _parse_exam_kind(raw: str, allowed: tuple[str, ...], type_label: str) -> str | None:
    s = _norm_header(raw)
    if not s:
        return None
    for code, aliases in EXAM_KIND_ALIASES.items():
        if s in aliases:
            if code not in allowed:
                raise _RowError(f"考试题型无效：{type_label}仅支持 {'/'.join(allowed)}")
            return code
    raise _RowError(f"未知考试题型：{(raw or '').strip()}")


def _parse_exam_level(raw: str) -> str | None:
    s = _norm_header(raw)
    if not s:
        return None
    for code, aliases in EXAM_LEVEL_ALIASES.items():
        if s in aliases:
            return code
    raise _RowError("考试级别无效，可选：" + "/".join(VOCAB_LEVEL_ORDER))


def _parse_bullets(raw: str) -> list[str] | None:
    """话题卡要点：仅按换行拆分（导出即换行连接，可逆）。

    不再按分号/竖线拆——要点本身含 ; 或 | 时按分隔符拆会改变结构，
    导出→导入往返不无损（评审 P2 实例：["where; why","who"] 变三条）。
    """
    s = (raw or "").strip()
    if not s:
        return None
    bullets = [p.strip() for p in s.splitlines()]
    bullets = [p for p in bullets if p]
    return bullets or None


def _parse_prep_seconds(raw: str) -> int | None:
    s = (raw or "").strip()
    if not s:
        return None
    try:
        n = int(s)
    except ValueError as exc:
        raise _RowError("准备秒数需为 10–180 之间的整数") from exc
    if not 10 <= n <= 180:
        raise _RowError("准备时间须在 10–180 秒之间")
    return n


# ── 行构建（校验 + 产出落库 payload；失败抛 _RowError） ──────────────


def _build_passage(cells: dict[str, str]) -> dict[str, Any]:
    title = (cells.get("title") or "").strip()
    if not title:
        raise _RowError("标题不能为空")
    if len(title) > 255:
        raise _RowError("标题过长（>255 字符）")
    topic = (cells.get("topic") or "").strip()
    if len(topic) > 100:
        raise _RowError("主题过长（>100 字符）")
    text = (cells.get("text") or "").strip()
    if not text:
        raise _RowError("正文不能为空")
    return {
        "title": title,
        "unit_title": (cells.get("unit_title") or "").strip() or None,
        "topic": topic,
        "cefr_band": _parse_band(cells.get("cefr_band") or ""),
        "text": text,
        "translation": _parse_optional_text(
            cells.get("translation") or "", max_len=1024, label="中文提示"
        ),
        "suggested_seconds": _parse_int(
            cells.get("suggested_seconds") or "",
            default=45,
            lo=10,
            hi=180,
            label="建议秒数",
        ),
        "audio_url": _parse_audio(cells.get("audio_url") or ""),
        "is_active": _parse_status(cells.get("is_active") or ""),
    }


def _build_repeat(cells: dict[str, str]) -> dict[str, Any]:
    text = (cells.get("text") or "").strip()
    if not text:
        raise _RowError("句子不能为空")
    if len(text) > 1024:
        raise _RowError("句子过长（>1024 字符）")
    exam_kind = _parse_exam_kind(
        cells.get("exam_kind") or "", allowed=EXAM_KIND_REPEAT, type_label="听句复述"
    )
    exam_level = _parse_exam_level(cells.get("exam_level") or "")
    err = _checked(lambda: validate_exam_fields("repeat", exam_kind, exam_level))
    if err:
        raise _RowError(err)
    return {
        "text": text,
        "translation": _parse_optional_text(
            cells.get("translation") or "", max_len=1024, label="中文提示"
        ),
        "passage_title": (cells.get("passage_title") or "").strip() or None,
        "suggested_seconds": _parse_int(
            cells.get("suggested_seconds") or "",
            default=8,
            lo=3,
            hi=60,
            label="作答秒数",
        ),
        "replay_limit": _parse_int(
            cells.get("replay_limit") or "", default=3, lo=0, hi=9, label="可听次数"
        ),
        "exam_kind": exam_kind,
        "exam_level": exam_level,
        "audio_url": _parse_audio(cells.get("audio_url") or ""),
    }


def _build_question(cells: dict[str, str]) -> dict[str, Any]:
    topic = (cells.get("topic") or "").strip()
    if not topic:
        raise _RowError("主题不能为空")
    if len(topic) > 100:
        raise _RowError("主题过长（>100 字符）")
    text = (cells.get("text") or "").strip()
    if not text:
        raise _RowError("题目不能为空")
    if len(text) > 512:
        raise _RowError("题目过长（>512 字符）")
    exam_kind = _parse_exam_kind(
        cells.get("exam_kind") or "", allowed=EXAM_KIND_QUESTION, type_label="情景问答"
    )
    exam_level = _parse_exam_level(cells.get("exam_level") or "")
    bullets = _parse_bullets(cells.get("cue_card_bullets") or "")
    prep_seconds = _parse_prep_seconds(cells.get("prep_seconds") or "")
    seconds = _parse_int(
        cells.get("suggested_seconds") or "",
        default=20,
        lo=10,
        hi=300,
        label="作答秒数",
    )
    err = _checked(
        lambda: validate_exam_fields(
            "question", exam_kind, exam_level, bullets, prep_seconds
        )
    ) or _checked(lambda: validate_question_suggested_seconds(seconds, exam_kind))
    if err:
        raise _RowError(err)
    return {
        "topic": topic,
        "text": text,
        "translation": _parse_optional_text(
            cells.get("translation") or "", max_len=1024, label="中文提示"
        ),
        "suggested_seconds": seconds,
        "exam_kind": exam_kind,
        "exam_level": exam_level,
        "cue_card_bullets": bullets,
        "prep_seconds": prep_seconds,
        "audio_url": _parse_audio(cells.get("audio_url") or ""),
    }


def _build_instruction(cells: dict[str, str]) -> dict[str, Any]:
    text = (cells.get("text") or "").strip()
    if not text:
        raise _RowError("说明文字不能为空")
    if len(text) > 2000:
        raise _RowError("说明文字过长（>2000 字符）")
    return {
        "title": _parse_optional_text(
            cells.get("title") or "", max_len=100, label="小标题"
        ),
        "text": text,
        "suggested_seconds": _parse_int(
            cells.get("suggested_seconds") or "", default=20, lo=5, hi=300, label="秒数"
        ),
    }


BUILDERS: dict[str, Any] = {
    TYPE_PASSAGE: _build_passage,
    TYPE_REPEAT: _build_repeat,
    TYPE_QUESTION: _build_question,
    TYPE_INSTRUCTION: _build_instruction,
}


# ── 文件解析（xlsx / CSV → 行列表 + 解析期问题） ────────────────────


@dataclass
class _RawRow:
    qtype: str
    line: int  # 文件内实际行号（表头为第 1 行，数据从 2 起）
    cells: dict[str, str]


@dataclass
class ParsedFile:
    rows: list[_RawRow]
    issues: list[QuestionBankImportIssue]  # 解析期问题（CSV 题型列无效等）


def _match_sheet_type(name: str) -> str | None:
    norm = _norm_header(name)
    for qtype, aliases in SHEET_NAME_ALIASES.items():
        if norm in aliases:
            return qtype
    return None


def _is_guide_sheet(name: str) -> bool:
    return _norm_header(name) in INSTRUCTION_SHEET_ALIASES


def parse_upload(filename: str, raw: bytes) -> ParsedFile:
    """解析上传文件 → 行列表（不做校验/查重）。文件级问题直接 422。"""
    name = (filename or "").casefold()
    if name.endswith(".xlsx") or raw[:2] == b"PK":
        parsed = _parse_xlsx(raw)
    else:
        # 其余一律按 CSV 尝试（utf-8-sig → gb18030 兜底：Excel 中文环境
        # 「另存为 CSV」默认 GBK，直接拒绝会把老师挡在门外）
        parsed = _parse_csv(raw)
    if len(parsed.rows) > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=422, detail=f"行数超过上限（{MAX_IMPORT_ROWS} 行），请分批导入"
        )
    if not parsed.rows and not parsed.issues:
        raise HTTPException(status_code=422, detail="文件中没有可导入的数据行")
    return parsed


def _parse_xlsx(raw: bytes) -> ParsedFile:
    """流式解析 xlsx，边读边限。

    行/列上限在读遍历**之前**按工作表声明范围（max_row/max_column）判定：
    只有格式的远端单元格（如整列设置格式）同样会扩大有效区——若读完再数
    数据行，一个 4.9KB 的文件就能让解析器遍历几十万个单元格（评审 P1）。
    无尺寸声明的合法文件 max_row/max_column 为 None（声明预检跳过），
    此时由遍历中的逐行实测兜底：行号上限 + 实际列数上限。
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=422, detail="Excel 文件无法解析，请用 Excel/WPS 重新保存后重试"
        ) from exc
    rows: list[_RawRow] = []
    issues: list[QuestionBankImportIssue] = []
    matched_any = False
    try:
        for sheet_name in workbook.sheetnames:
            if _is_guide_sheet(sheet_name):
                continue  # 说明页按名字跳过，不读内容
            sheet = workbook[sheet_name]
            max_row = sheet.max_row or 0
            max_column = sheet.max_column or 0
            if max_column > MAX_SHEET_COLUMNS:
                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"工作表「{sheet_name}」有 {max_column} 列，超过 "
                        f"{MAX_SHEET_COLUMNS} 列上限，请使用标准模板的列"
                    ),
                )
            if max_row > MAX_SHEET_ROWS:
                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"工作表「{sheet_name}」有效区共 {max_row} 行，超过 "
                        f"{MAX_IMPORT_ROWS} 行数据上限；若因整列设置格式导致，"
                        "请清除多余格式或把数据复制到新模板"
                    ),
                )
            header_done = False
            col_map: dict[int, str] = {}
            qtype: str | None = None
            # max_row=0（无尺寸声明）时不能下传 max_row——会一个单元格都不产出
            iter_kwargs: dict[str, Any] = {"values_only": True}
            if max_row:
                iter_kwargs["max_row"] = max_row
            for r_idx, row in enumerate(sheet.iter_rows(**iter_kwargs), start=1):
                # 防御声明范围偏小/缺失的畸形文件：实际行数超界立即拒绝
                if r_idx > MAX_SHEET_ROWS:
                    raise HTTPException(
                        status_code=422,
                        detail=f"行数超过上限（{MAX_IMPORT_ROWS} 行），请分批导入",
                    )
                # 无尺寸声明（max_column=None）时声明预检失效：逐行实测列数
                if len(row) > MAX_SHEET_COLUMNS:
                    raise HTTPException(
                        status_code=422,
                        detail=(
                            f"工作表「{sheet_name}」有 {len(row)} 列，超过 "
                            f"{MAX_SHEET_COLUMNS} 列上限，请使用标准模板的列"
                        ),
                    )
                row_cells = tuple(_cell_str(c) for c in row)
                if not any(row_cells):
                    continue
                if not header_done:
                    header_done = True
                    qtype = _match_sheet_type(sheet_name)
                    if qtype is None:
                        raise HTTPException(
                            status_code=422,
                            detail=(
                                f"工作表「{sheet_name}」无法识别题型，"
                                "请勿重命名模板工作表"
                            ),
                        )
                    matched_any = True
                    col_map = _map_header(row_cells)
                    _require_columns(qtype, col_map, sheet_name)
                    continue
                if len(rows) + 1 > MAX_IMPORT_ROWS:
                    raise HTTPException(
                        status_code=422,
                        detail=f"行数超过上限（{MAX_IMPORT_ROWS} 行），请分批导入",
                    )
                rows.append(
                    _RawRow(
                        qtype=qtype or "",
                        line=r_idx,
                        cells=_row_cells(col_map, row_cells),
                    )
                )
    finally:
        workbook.close()
    if not matched_any:
        raise HTTPException(
            status_code=422,
            detail="未找到可识别的工作表（需包含：文章朗读/听句复述/情景问答/题目说明）",
        )
    return ParsedFile(rows=rows, issues=issues)


def _parse_csv(raw: bytes) -> ParsedFile:
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=422, detail="文件必须是 UTF-8 编码")
    reader = csv.reader(io.StringIO(text))
    header: list[str] | None = None
    col_map: dict[int, str] = {}
    type_col: int | None = None
    fields_present: set[str] = set()
    rows: list[_RawRow] = []
    issues: list[QuestionBankImportIssue] = []
    checked_types: set[str] = set()
    any_content = False
    # 数据行计数包含无效行（题型空/无效）——只数成功识别的行会让
    # 2001 行全无效的文件绕过行数上限（评审 P2）
    data_rows = 0
    try:
        for r_idx, raw_row in enumerate(reader, start=1):
            # 还原导出时的公式防护前缀（见 _csv_unguard）：往返必须无损
            row_cells = tuple(_csv_unguard(c.strip()) for c in raw_row)
            if not any(row_cells):
                continue
            any_content = True
            if header is None:
                header = list(row_cells)
                col_map = _map_header(header)
                type_col = next(
                    (c for c, f in col_map.items() if f == "__type__"), None
                )
                if type_col is None:
                    raise HTTPException(status_code=422, detail="CSV 缺少题型列")
                fields_present = set(col_map.values())
                continue
            data_rows += 1
            if data_rows > MAX_IMPORT_ROWS:
                raise HTTPException(
                    status_code=422,
                    detail=f"行数超过上限（{MAX_IMPORT_ROWS} 行），请分批导入",
                )
            assert type_col is not None
            type_cell = row_cells[type_col] if type_col < len(row_cells) else ""
            if not type_cell:
                issues.append(
                    QuestionBankImportIssue(type="-", line=r_idx, reason="题型不能为空")
                )
                continue
            qtype = _match_sheet_type(type_cell)
            if qtype is None:
                issues.append(
                    QuestionBankImportIssue(
                        type="-", line=r_idx, reason=f"题型列的值无效：{type_cell}"
                    )
                )
                continue
            if qtype not in checked_types:
                checked_types.add(qtype)
                _require_columns_csv(qtype, fields_present)
            cells = _row_cells(
                {c: f for c, f in col_map.items() if f != "__type__"}, row_cells
            )
            rows.append(_RawRow(qtype=qtype, line=r_idx, cells=cells))
    except csv.Error as exc:
        # field_size_limit 已抬高到上传上限之上，此处理论不可达——保留
        # 防御：任何 csv 模块级错误都以 422 呈现，绝不 500
        raise HTTPException(status_code=422, detail="CSV 字段过大，无法解析") from exc
    if not any_content:
        raise HTTPException(status_code=422, detail="文件为空")
    return ParsedFile(rows=rows, issues=issues)


def _map_header(header_cells: tuple[str, ...] | list[str]) -> dict[int, str]:
    """列号 → canonical 字段（未识别列忽略）。"""
    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_cells):
        mapped = HEADER_LOOKUP.get(_norm_header(cell))
        if mapped is not None:
            col_map[idx] = mapped
    return col_map


def _zh_of(qtype: str, fieldname: str) -> str:
    return next(c.zh for c in COLUMNS_BY_TYPE[qtype] if c.field == fieldname)


def _require_columns(qtype: str, col_map: dict[int, str], sheet_name: str) -> None:
    present = set(col_map.values())
    missing = [f for f in REQUIRED_FIELDS[qtype] if f not in present]
    if missing:
        names = "/".join(_zh_of(qtype, f) for f in missing)
        raise HTTPException(
            status_code=422, detail=f"工作表「{sheet_name}」缺少必填列：{names}"
        )


def _require_columns_csv(qtype: str, fields_present: set[str]) -> None:
    missing = [f for f in REQUIRED_FIELDS[qtype] if f not in fields_present]
    if missing:
        names = "/".join(_zh_of(qtype, f) for f in missing)
        raise HTTPException(
            status_code=422, detail=f"CSV 表头缺少{TYPE_LABELS[qtype]}必填列：{names}"
        )


def _row_cells(
    col_map: dict[int, str], row_cells: tuple[str, ...] | list[str]
) -> dict[str, str]:
    cells: dict[str, str] = {}
    for idx, fieldname in col_map.items():
        if idx < len(row_cells):
            cells[fieldname] = row_cells[idx]
    return cells


# ── 库内参照与查重计划 ───────────────────────────────────────────────


@dataclass
class _Ref:
    """查重与兜底解析需要的库内参照（一次性加载；题库规模下全表可控）。"""

    units_by_norm: dict[str, list[Unit]]
    root_passages_by_norm: dict[str, list[Passage]]
    passage_exact: dict[tuple[str, str], list[tuple[Any, ...]]]
    passage_texts: set[str]
    # 主题匹配：精确优先。规范化字典允许 Travel/travel 并存（后者不覆盖
    # 前者）；规范化命中多个时由调用方报行级错误，绝不静默选一个
    scenarios_by_exact: dict[str, Scenario]
    scenarios_by_norm: dict[str, list[Scenario]]
    question_exact: dict[tuple[str, str], list[tuple[Any, ...]]]
    question_texts: dict[str, set[str]]
    sentence_configs: dict[str, list[tuple[Any, ...]]]
    instruction_configs: dict[str, list[tuple[Any, ...]]]
    existing_slugs: set[str]
    max_question_order: dict[uuid.UUID, int]
    max_sentence_order_by_passage: dict[uuid.UUID, int]
    max_sentence_order_global: int


def _load_reference(session: Session) -> _Ref:
    units = session.exec(select(Unit)).all()
    root_passages = session.exec(
        select(Passage).where(col(Passage.parent_passage_id).is_(None))
    ).all()
    all_passages = session.exec(select(Passage)).all()
    scenarios = session.exec(select(Scenario)).all()
    questions = session.exec(select(ScenarioQuestion)).all()
    sentences = session.exec(select(RepeatSentence)).all()
    instructions = session.exec(select(Instruction)).all()

    units_by_norm: dict[str, list[Unit]] = {}
    for unit in units:
        units_by_norm.setdefault(_norm_text(unit.title), []).append(unit)

    root_passages_by_norm: dict[str, list[Passage]] = {}
    for passage in root_passages:
        root_passages_by_norm.setdefault(_norm_text(passage.title), []).append(passage)

    unit_titles = {u.id: u.title for u in units}
    passage_titles = {p.id: p.title for p in all_passages}
    passage_exact: dict[tuple[str, str], list[tuple[Any, ...]]] = {}
    passage_texts: set[str] = set()
    for passage in root_passages:
        key = (_norm_text(passage.title), _norm_text(passage.text))
        cfg = (
            _norm_text(passage.translation or ""),
            passage.cefr_band,
            passage.suggested_seconds,
            passage.is_active,
            _norm_text(unit_titles.get(passage.unit_id, "")) if passage.unit_id else "",
            _norm_text(passage.topic or ""),
            passage.audio_url or "",
        )
        passage_exact.setdefault(key, []).append(cfg)
        passage_texts.add(key[1])

    scenarios_by_exact = {s.topic: s for s in scenarios}
    scenarios_by_norm: dict[str, list[Scenario]] = {}
    for s in scenarios:
        scenarios_by_norm.setdefault(_norm_text(s.topic), []).append(s)
    topic_by_scenario = {s.id: s.topic for s in scenarios}
    question_exact: dict[tuple[str, str], list[tuple[Any, ...]]] = {}
    question_texts: dict[str, set[str]] = {}
    for question in questions:
        topic = topic_by_scenario.get(question.scenario_id, "")
        key = (_norm_text(topic), _norm_text(question.text))
        cfg = (
            _norm_text(question.translation or ""),
            question.suggested_seconds,
            question.exam_kind or "",
            question.exam_level or "",
            tuple(question.cue_card_bullets or ()),
            question.prep_seconds or 0,
            question.audio_url or "",
        )
        question_exact.setdefault(key, []).append(cfg)
        question_texts.setdefault(_norm_text(question.text), set()).add(
            _norm_text(topic)
        )

    sentence_configs: dict[str, list[tuple[Any, ...]]] = {}
    for sentence in sentences:
        cfg = (
            _norm_text(sentence.translation or ""),
            sentence.suggested_seconds,
            sentence.replay_limit,
            sentence.exam_kind or "",
            sentence.exam_level or "",
            _norm_text(passage_titles.get(sentence.passage_id, ""))
            if sentence.passage_id
            else "",
            sentence.audio_url or "",
        )
        sentence_configs.setdefault(_norm_text(sentence.text), []).append(cfg)

    instruction_configs: dict[str, list[tuple[Any, ...]]] = {}
    for ins in instructions:
        cfg = (_norm_text(ins.title or ""), ins.suggested_seconds)
        instruction_configs.setdefault(_norm_text(ins.text), []).append(cfg)

    max_question_order: dict[uuid.UUID, int] = {}
    for question in questions:
        max_question_order[question.scenario_id] = max(
            max_question_order.get(question.scenario_id, 0), question.order_index
        )
    max_by_passage: dict[uuid.UUID, int] = {}
    global_max = 0
    for sentence in sentences:
        global_max = max(global_max, sentence.order_index)
        if sentence.passage_id is not None:
            max_by_passage[sentence.passage_id] = max(
                max_by_passage.get(sentence.passage_id, 0), sentence.order_index
            )

    return _Ref(
        units_by_norm=units_by_norm,
        root_passages_by_norm=root_passages_by_norm,
        passage_exact=passage_exact,
        passage_texts=passage_texts,
        scenarios_by_exact=scenarios_by_exact,
        scenarios_by_norm=scenarios_by_norm,
        question_exact=question_exact,
        question_texts=question_texts,
        sentence_configs=sentence_configs,
        instruction_configs=instruction_configs,
        existing_slugs={p.slug for p in all_passages},
        max_question_order=max_question_order,
        max_sentence_order_by_passage=max_by_passage,
        max_sentence_order_global=global_max,
    )


@dataclass
class _TypeCounts:
    valid: int = 0
    exact_dup: int = 0
    suspected: int = 0
    invalid: int = 0
    missing_audio: int = 0


@dataclass
class _CreateItem:
    line: int
    payload: dict[str, Any]
    # 听句复述挂靠：命中库内篇目 → passage_id；命中本次导入篇目 → 计划标题；
    # 都没有 → None（独立句）
    passage_id: uuid.UUID | None = None
    planned_passage_title: str | None = None
    # 情景问答主题：命中库内 → scenario_id；否则 new_topic（确认时新建）
    scenario_id: uuid.UUID | None = None
    new_topic: str | None = None


@dataclass
class QuestionBankPlan:
    creates: dict[str, list[_CreateItem]] = field(default_factory=dict)
    counts: dict[str, _TypeCounts] = field(default_factory=dict)
    issues: list[QuestionBankImportIssue] = field(default_factory=list)
    suspected_rows: list[QuestionBankImportIssue] = field(default_factory=list)
    new_topics: list[str] = field(default_factory=list)
    unit_fallback_count: int = 0
    standalone_sentence_count: int = 0
    # apply 阶段需要的库内状态（与 plan 同一快照）
    existing_slugs: set[str] = field(default_factory=set)
    max_question_order: dict[uuid.UUID, int] = field(default_factory=dict)
    max_sentence_order_by_passage: dict[uuid.UUID, int] = field(default_factory=dict)
    max_sentence_order_global: int = 0


def plan_import(
    session: Session, parsed: ParsedFile, include_suspected: bool = False
) -> QuestionBankPlan:
    """校验 + 查重 + 兜底解析，产出创建计划（不写库）。"""
    ref = _load_reference(session)
    plan = QuestionBankPlan(
        existing_slugs=ref.existing_slugs,
        max_question_order=ref.max_question_order,
        max_sentence_order_by_passage=ref.max_sentence_order_by_passage,
        max_sentence_order_global=ref.max_sentence_order_global,
    )
    for qtype in QUESTION_BANK_TYPE_ORDER:
        plan.creates[qtype] = []
        plan.counts[qtype] = _TypeCounts()
    plan.issues.extend(parsed.issues)

    by_type: dict[str, list[_RawRow]] = {t: [] for t in plan.creates}
    for row in parsed.rows:
        by_type.setdefault(row.qtype, []).append(row)

    _plan_passages(plan, ref, by_type[TYPE_PASSAGE], include_suspected)
    _plan_repeats(plan, ref, by_type[TYPE_REPEAT], include_suspected)
    _plan_questions(plan, ref, by_type[TYPE_QUESTION], include_suspected)
    _plan_instructions(plan, ref, by_type[TYPE_INSTRUCTION], include_suspected)
    return plan


def _issue(qtype: str, line: int, reason: str) -> QuestionBankImportIssue:
    return QuestionBankImportIssue(type=qtype, line=line, reason=reason)


def _plan_passages(
    plan: QuestionBankPlan, ref: _Ref, rows: list[_RawRow], include_suspected: bool
) -> None:
    counts = plan.counts[TYPE_PASSAGE]
    planned_exact: dict[tuple[str, str], list[tuple[Any, ...]]] = {}
    planned_texts: set[str] = set()
    for row in rows:
        try:
            payload = _build_passage(row.cells)
        except _RowError as exc:
            counts.invalid += 1
            plan.issues.append(_issue(TYPE_PASSAGE, row.line, exc.reason))
            continue
        n_title, n_text = _norm_text(payload["title"]), _norm_text(payload["text"])
        key = (n_title, n_text)
        # 单元解析必须先于查重：落库时挂单元的篇目 topic 会被派生为单元主题
        # （crud.derive_passage_topic），查重配置若用原始 topic，同一文件重传
        # 会因「落库值≠导入值」被误判疑似重复（评审 P2 实例）
        unit_title_resolved = ""
        topic_resolved = payload["topic"]
        if payload["unit_title"] is not None:
            candidates = ref.units_by_norm.get(_norm_text(payload["unit_title"]), [])
            if len(candidates) > 1:
                counts.invalid += 1
                plan.issues.append(
                    _issue(
                        TYPE_PASSAGE,
                        row.line,
                        "单元匹配到多个同名单元，请改为精确的单元标题",
                    )
                )
                continue
            if len(candidates) == 1:
                payload["unit_id"] = candidates[0].id
                unit_title_resolved = candidates[0].title
                topic_resolved = candidates[0].topic
            else:
                plan.unit_fallback_count += 1
        cfg = (
            _norm_text(payload["translation"] or ""),
            payload["cefr_band"],
            payload["suggested_seconds"],
            payload["is_active"],
            _norm_text(unit_title_resolved),
            _norm_text(topic_resolved),
            payload["audio_url"] or "",
        )
        db_cfgs = ref.passage_exact.get(key, [])
        planned_cfgs = planned_exact.get(key, [])
        if db_cfgs or planned_cfgs:
            if cfg in db_cfgs or cfg in planned_cfgs:
                counts.exact_dup += 1
                continue
            counts.suspected += 1
            plan.suspected_rows.append(
                _issue(
                    TYPE_PASSAGE,
                    row.line,
                    "疑似重复：同名同正文篇目已存在，但难度/秒数等配置不同",
                )
            )
            if not include_suspected:
                continue
        elif n_text in ref.passage_texts or n_text in planned_texts:
            # 批内正文集合也要查：同批两篇不同标题、相同正文 → 第二篇疑似
            counts.suspected += 1
            plan.suspected_rows.append(
                _issue(TYPE_PASSAGE, row.line, "疑似重复：相同正文已存在于其他篇目")
            )
            if not include_suspected:
                continue
        if not payload["audio_url"]:
            counts.missing_audio += 1
        counts.valid += 1
        plan.creates[TYPE_PASSAGE].append(_CreateItem(line=row.line, payload=payload))
        planned_exact.setdefault(key, []).append(cfg)
        planned_texts.add(n_text)


def _plan_repeats(
    plan: QuestionBankPlan, ref: _Ref, rows: list[_RawRow], include_suspected: bool
) -> None:
    counts = plan.counts[TYPE_REPEAT]
    planned_configs: dict[str, list[tuple[Any, ...]]] = {}
    # 本批计划新建的篇目标题（供挂靠；同批多篇同名 → 歧义报错）
    planned_titles: dict[str, int] = {}
    for item in plan.creates[TYPE_PASSAGE]:
        n_title = _norm_text(item.payload["title"])
        planned_titles[n_title] = planned_titles.get(n_title, 0) + 1
    for row in rows:
        try:
            payload = _build_repeat(row.cells)
        except _RowError as exc:
            counts.invalid += 1
            plan.issues.append(_issue(TYPE_REPEAT, row.line, exc.reason))
            continue
        n_text = _norm_text(payload["text"])
        # 挂靠解析必须先于查重：cfg 的挂靠分量要用**落库后的最终关联**——
        # 命中库内/本批篇目 → 规范化篇目标题；找不到而降级独立句 → 空串
        # （与库内独立句的 cfg 一致）。否则「挂靠 Ghost」首次导入成独立句后，
        # 重传同文件会因 cfg 不匹配被误判疑似、勾选后建出重复独立句（评审 P2）
        passage_id: uuid.UUID | None = None
        planned_title: str | None = None
        linkage_norm = ""
        if payload["passage_title"] is not None:
            n_title = _norm_text(payload["passage_title"])
            existing_hits = len(ref.root_passages_by_norm.get(n_title, []))
            planned_hits = planned_titles.get(n_title, 0)
            if existing_hits + planned_hits > 1:
                counts.invalid += 1
                plan.issues.append(
                    _issue(
                        TYPE_REPEAT,
                        row.line,
                        "挂靠篇目匹配到多个同名篇目，请先重命名区分",
                    )
                )
                continue
            if existing_hits == 1:
                passage_id = ref.root_passages_by_norm[n_title][0].id
                linkage_norm = n_title
            elif planned_hits == 1:
                planned_title = n_title
                linkage_norm = n_title
            else:
                plan.standalone_sentence_count += 1
        cfg = (
            _norm_text(payload["translation"] or ""),
            payload["suggested_seconds"],
            payload["replay_limit"],
            payload["exam_kind"] or "",
            payload["exam_level"] or "",
            linkage_norm,
            payload["audio_url"] or "",
        )
        db_cfgs = ref.sentence_configs.get(n_text, [])
        planned_cfgs = planned_configs.get(n_text, [])
        if db_cfgs or planned_cfgs:
            if cfg in db_cfgs or cfg in planned_cfgs:
                counts.exact_dup += 1
                continue
            counts.suspected += 1
            plan.suspected_rows.append(
                _issue(TYPE_REPEAT, row.line, "疑似重复：相同句子已存在但配置不同")
            )
            if not include_suspected:
                continue
        if not payload["audio_url"]:
            counts.missing_audio += 1
        counts.valid += 1
        plan.creates[TYPE_REPEAT].append(
            _CreateItem(
                line=row.line,
                payload=payload,
                passage_id=passage_id,
                planned_passage_title=planned_title,
            )
        )
        planned_configs.setdefault(n_text, []).append(cfg)


def _plan_questions(
    plan: QuestionBankPlan, ref: _Ref, rows: list[_RawRow], include_suspected: bool
) -> None:
    counts = plan.counts[TYPE_QUESTION]
    planned_exact: dict[tuple[str, str], list[tuple[Any, ...]]] = {}
    planned_texts: dict[str, set[str]] = {}
    new_topic_map: dict[str, str] = {}
    for row in rows:
        try:
            payload = _build_question(row.cells)
        except _RowError as exc:
            counts.invalid += 1
            plan.issues.append(_issue(TYPE_QUESTION, row.line, exc.reason))
            continue
        n_topic, n_text = _norm_text(payload["topic"]), _norm_text(payload["text"])
        key = (n_topic, n_text)
        cfg = (
            _norm_text(payload["translation"] or ""),
            payload["suggested_seconds"],
            payload["exam_kind"] or "",
            payload["exam_level"] or "",
            tuple(payload["cue_card_bullets"] or ()),
            payload["prep_seconds"] or 0,
            payload["audio_url"] or "",
        )
        db_cfgs = ref.question_exact.get(key, [])
        planned_cfgs = planned_exact.get(key, [])
        tier: str | None = None
        if db_cfgs or planned_cfgs:
            tier = "exact" if cfg in db_cfgs or cfg in planned_cfgs else "same_topic"
        elif n_text in ref.question_texts or n_text in planned_texts:
            tier = "cross_topic"
        if tier == "exact":
            counts.exact_dup += 1
            continue
        if tier is not None:
            counts.suspected += 1
            reason = (
                "疑似重复：该主题下已有相同题目但配置不同"
                if tier == "same_topic"
                else "疑似重复：该题目已存在于其他主题"
            )
            plan.suspected_rows.append(_issue(TYPE_QUESTION, row.line, reason))
            if not include_suspected:
                continue
        # 主题解析精确优先：库内允许 Travel/travel 并存，规范化字典只作
        # 兜底；规范化命中多个时报行级错误，绝不静默挑一个（评审 P2）
        scenario = ref.scenarios_by_exact.get(payload["topic"])
        if scenario is None:
            candidates = ref.scenarios_by_norm.get(n_topic, [])
            if len(candidates) > 1:
                counts.invalid += 1
                plan.issues.append(
                    _issue(
                        TYPE_QUESTION,
                        row.line,
                        "主题匹配到多个大小写不同的同名主题，请使用与题库完全一致的主题名",
                    )
                )
                continue
            if len(candidates) == 1:
                scenario = candidates[0]
        if scenario is None and n_topic not in new_topic_map:
            new_topic_map[n_topic] = payload["topic"]
            plan.new_topics.append(payload["topic"])
        if not payload["audio_url"]:
            counts.missing_audio += 1
        counts.valid += 1
        plan.creates[TYPE_QUESTION].append(
            _CreateItem(
                line=row.line,
                payload=payload,
                scenario_id=scenario.id if scenario is not None else None,
                new_topic=None if scenario is not None else new_topic_map[n_topic],
            )
        )
        planned_exact.setdefault(key, []).append(cfg)
        planned_texts.setdefault(n_text, set()).add(n_topic)


def _plan_instructions(
    plan: QuestionBankPlan, ref: _Ref, rows: list[_RawRow], include_suspected: bool
) -> None:
    counts = plan.counts[TYPE_INSTRUCTION]
    planned_configs: dict[str, list[tuple[Any, ...]]] = {}
    for row in rows:
        try:
            payload = _build_instruction(row.cells)
        except _RowError as exc:
            counts.invalid += 1
            plan.issues.append(_issue(TYPE_INSTRUCTION, row.line, exc.reason))
            continue
        n_text = _norm_text(payload["text"])
        cfg = (_norm_text(payload["title"] or ""), payload["suggested_seconds"])
        db_cfgs = ref.instruction_configs.get(n_text, [])
        planned_cfgs = planned_configs.get(n_text, [])
        if db_cfgs or planned_cfgs:
            if cfg in db_cfgs or cfg in planned_cfgs:
                counts.exact_dup += 1
                continue
            counts.suspected += 1
            plan.suspected_rows.append(
                _issue(
                    TYPE_INSTRUCTION,
                    row.line,
                    "疑似重复：相同说明文字已存在但标题/秒数不同",
                )
            )
            if not include_suspected:
                continue
        counts.valid += 1
        plan.creates[TYPE_INSTRUCTION].append(
            _CreateItem(line=row.line, payload=payload)
        )
        planned_configs.setdefault(n_text, []).append(cfg)


# ── 预览与落库 ───────────────────────────────────────────────────────


def preview_import(session: Session, parsed: ParsedFile) -> QuestionBankImportPreview:
    plan = plan_import(session, parsed, include_suspected=False)
    sheets = []
    total_valid = 0
    for qtype in QUESTION_BANK_TYPE_ORDER:
        counts = plan.counts[qtype]
        if counts.valid + counts.exact_dup + counts.suspected + counts.invalid == 0:
            continue
        sheets.append(
            QuestionBankTypeCount(
                type=qtype,
                valid_count=counts.valid,
                exact_duplicate_count=counts.exact_dup,
                suspected_duplicate_count=counts.suspected,
                missing_audio_count=counts.missing_audio,
            )
        )
        total_valid += counts.valid
    return QuestionBankImportPreview(
        sheets=sheets,
        issues=plan.issues[:MAX_ISSUES],
        issues_truncated=len(plan.issues) > MAX_ISSUES,
        suspected=plan.suspected_rows[:MAX_SUSPECTED_LISTED],
        suspected_truncated=len(plan.suspected_rows) > MAX_SUSPECTED_LISTED,
        new_topics=plan.new_topics,
        unit_fallback_count=plan.unit_fallback_count,
        standalone_sentence_count=plan.standalone_sentence_count,
        total_valid=total_valid,
    )


def _commit(session: Session) -> None:
    """提交封装：测试可注入失败验证整体回滚（vocab_levels 同款）。"""
    session.commit()


def apply_import(
    session: Session, parsed: ParsedFile, include_suspected: bool
) -> QuestionBankImportResult:
    """确认导入：重新解析校验后单事务落库；任何 DB 异常整体回滚。"""
    plan = plan_import(session, parsed, include_suspected)
    try:
        created_scenarios: dict[str, Scenario] = {}
        for topic in plan.new_topics:
            scenario = Scenario(topic=topic, is_active=True)
            session.add(scenario)
            created_scenarios[_norm_text(topic)] = scenario
        session.flush()  # 分配 scenario.id

        used_slugs = set(plan.existing_slugs)
        created_passages: dict[str, Passage] = {}
        for item in plan.creates[TYPE_PASSAGE]:
            payload = item.payload
            passage = Passage(
                title=payload["title"],
                topic=crud.derive_passage_topic(
                    session, payload.get("unit_id"), payload["topic"]
                ),
                cefr_band=payload["cefr_band"],
                text=payload["text"],
                translation=payload["translation"],
                audio_url=payload["audio_url"],
                suggested_seconds=payload["suggested_seconds"],
                is_active=payload["is_active"],
                slug=_generate_slug(payload["title"], used_slugs),
                unit_id=payload.get("unit_id"),
            )
            session.add(passage)
            created_passages[_norm_text(passage.title)] = passage
        session.flush()  # 分配 passage.id，供复述句挂靠

        per_passage_order = dict(plan.max_sentence_order_by_passage)
        global_order = plan.max_sentence_order_global
        for item in plan.creates[TYPE_REPEAT]:
            payload = item.payload
            passage_id = item.passage_id
            if passage_id is None and item.planned_passage_title is not None:
                planned = created_passages.get(item.planned_passage_title)
                passage_id = planned.id if planned is not None else None
            if passage_id is not None:
                order = per_passage_order.get(passage_id, 0) + 1
                per_passage_order[passage_id] = order
            else:
                global_order += 1
                order = global_order
            session.add(
                RepeatSentence(
                    passage_id=passage_id,
                    order_index=order,
                    text=payload["text"],
                    translation=payload["translation"],
                    audio_url=payload["audio_url"],
                    suggested_seconds=payload["suggested_seconds"],
                    replay_limit=payload["replay_limit"],
                    exam_kind=payload["exam_kind"],
                    exam_level=payload["exam_level"],
                )
            )

        question_order = dict(plan.max_question_order)
        for item in plan.creates[TYPE_QUESTION]:
            payload = item.payload
            if item.scenario_id is not None:
                scenario_id = item.scenario_id
            else:
                scenario = created_scenarios[_norm_text(item.new_topic or "")]
                scenario_id = scenario.id
            order = question_order.get(scenario_id, 0) + 1
            question_order[scenario_id] = order
            session.add(
                ScenarioQuestion(
                    scenario_id=scenario_id,
                    band="B1",
                    order_index=order,
                    text=payload["text"],
                    translation=payload["translation"],
                    audio_url=payload["audio_url"],
                    suggested_seconds=payload["suggested_seconds"],
                    exam_kind=payload["exam_kind"],
                    exam_level=payload["exam_level"],
                    cue_card_bullets=payload["cue_card_bullets"],
                    prep_seconds=payload["prep_seconds"],
                )
            )

        for item in plan.creates[TYPE_INSTRUCTION]:
            payload = item.payload
            session.add(
                Instruction(
                    title=payload["title"],
                    text=payload["text"],
                    suggested_seconds=payload["suggested_seconds"],
                )
            )
        _commit(session)
    except HTTPException:
        raise
    except Exception as exc:
        session.rollback()
        raise HTTPException(
            status_code=500, detail="导入失败，已整体回滚，请重试"
        ) from exc

    sheets = []
    missing_total = 0
    for qtype in QUESTION_BANK_TYPE_ORDER:
        counts = plan.counts[qtype]
        if not any(
            (
                plan.creates[qtype],
                counts.exact_dup,
                counts.suspected if not include_suspected else 0,
                counts.invalid,
            )
        ):
            continue
        missing_total += counts.missing_audio
        sheets.append(
            QuestionBankImportTypeResult(
                type=qtype,
                created=len(plan.creates[qtype]),
                skipped_duplicate=counts.exact_dup,
                skipped_suspected=0 if include_suspected else counts.suspected,
                skipped_invalid=counts.invalid,
                missing_audio_count=counts.missing_audio,
            )
        )
    return QuestionBankImportResult(
        sheets=sheets,
        created_topics=plan.new_topics,
        issues=plan.issues[:MAX_ISSUES],
        missing_audio_total=missing_total,
    )


def _generate_slug(title: str, used: set[str]) -> str:
    """标题 → 唯一 slug（与 crud.unique_passage_slug 同规则，批内去重免逐行查询）。"""
    base = crud.slugify_title(title)
    slug = base
    for i in range(2, 500):
        if slug not in used:
            break
        suffix = f"-{i}"
        slug = f"{base[: 100 - len(suffix)]}{suffix}"
    else:
        slug = f"{base[:90]}-{uuid.uuid4().hex[:8]}"
    used.add(slug)
    return slug


# ── 模板生成 ─────────────────────────────────────────────────────────

GUIDE_LINES = (
    "题库批量导入模板 · Question Bank Import Template",
    "",
    "使用步骤 Steps：",
    "1. 按题型填写对应工作表（每个工作表只放一种题型）Fill one sheet per question type",
    "2. 上传到「题目库 → 批量导入」预览 Upload in Question Bank → Bulk import and preview",
    "3. 确认无误后点「确认导入」 Confirm to import",
    "",
    "必填列 Required columns（其余列留空即用默认值；灰色示例行可删除）Leave other columns blank for defaults; example rows can be deleted：",
    "  文章朗读 Read Aloud：标题、正文 Title, Text",
    "  听句复述 Listen & Repeat：英文句子 Sentence",
    "  情景问答 Scenario Q&A：主题、英文题目 Topic, Question",
    "  题目说明 Instruction：说明文字 Text",
    "",
    "下拉选项 Dropdowns：难度/考试题型/考试级别/可听次数/状态请从下拉列表选择",
    "CSV 导入：单文件用「题型」列区分题型（值同工作表名）For CSV, use the Type column",
    "",
    "重复处理 Duplicates：完全重复自动跳过；疑似重复（内容相同、配置不同）导入时会请你确认",
    "Exact duplicates are skipped automatically; near duplicates need your confirmation",
    "导入不会修改已有题目 Import never modifies existing questions",
    "",
    "音频 Audio：先导入文字，再在题目编辑里上传音频文件；「音频链接」列只填本站链接",
    "Import text first, then upload audio files in the question editor; the Audio column only accepts this site's links",
    "",
    "上限 Limits：导入 2000 行 / 5MB；导出同为 2000 行（超出请按题型/主题分批导出）",
    "Limits: import 2000 rows / 5MB; exports share the same 2000-row cap (split by type/topic when larger)",
    "编码 Encoding：Excel(.xlsx) 直接上传；CSV 需 UTF-8（推荐用本站下载的 CSV 模板）",
    "",
    "注意：本模板用于题目内容交换，不是完整备份——拆句朗读、逐题计时、题目说明的考试音频等高级配置不入模板，导入后如需请在题目编辑页设置",
    "Note: this template exchanges question content, not a full backup — advanced settings (per-sentence split, per-item timing, exam audio on instructions) must be reconfigured in the editor after import",
)

TEMPLATE_EXAMPLES: dict[str, list[list[str]]] = {
    TYPE_PASSAGE: [
        [
            "My Morning Routine",
            "Unit 1",
            "Daily Life",
            "B1",
            "I usually wake up at six o'clock. After breakfast, I walk to school with my neighbour.",
            "我通常六点起床。早饭后我和邻居一起走路去上学。",
            "45",
            "",
            "启用",
        ],
    ],
    TYPE_REPEAT: [
        [
            "The library opens at nine o'clock.",
            "图书馆九点开门。",
            "",
            "8",
            "3",
            "",
            "",
            "",
        ],
    ],
    TYPE_QUESTION: [
        [
            "Travel",
            "Do you prefer travelling alone or with friends? Why?",
            "你喜欢独自旅行还是和朋友一起？为什么？",
            "20",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "Travel",
            "Describe a place you visited that you would like to visit again.",
            "描述一个你去过、还想再去的地方。",
            "120",
            "IELTS · Part 2",
            "PET",
            "what the place is\nwhen you went there\nwhy you like it",
            "60",
            "",
        ],
    ],
    TYPE_INSTRUCTION: [
        [
            "Part B · 开始",
            "In this part, you will talk about the topic on the card. You will have one minute to prepare.",
            "60",
        ],
    ],
}

TEMPLATE_DROPDOWNS: dict[str, dict[str, list[str]]] = {
    TYPE_PASSAGE: {
        "cefr_band": ["A2", "B1", "B2"],
        "is_active": ["启用", "停用"],
    },
    TYPE_REPEAT: {
        "replay_limit": [str(i) for i in range(10)],
        "exam_kind": ["TOEFL · 听令复述"],
        "exam_level": list(VOCAB_LEVEL_ORDER),
    },
    TYPE_QUESTION: {
        "exam_kind": [
            "TOEFL · 情景访谈",
            "IELTS · Part 1",
            "IELTS · Part 2",
            "IELTS · Part 3",
        ],
        "exam_level": list(VOCAB_LEVEL_ORDER),
    },
    TYPE_INSTRUCTION: {},
}

TEMPLATE_WIDTHS: dict[str, dict[str, int]] = {
    TYPE_PASSAGE: {
        "title": 26,
        "unit_title": 14,
        "topic": 16,
        "cefr_band": 9,
        "text": 60,
        "translation": 34,
        "suggested_seconds": 11,
        "audio_url": 30,
        "is_active": 9,
    },
    TYPE_REPEAT: {
        "text": 50,
        "translation": 30,
        "passage_title": 20,
        "suggested_seconds": 11,
        "replay_limit": 11,
        "exam_kind": 18,
        "exam_level": 14,
        "audio_url": 30,
    },
    TYPE_QUESTION: {
        "topic": 16,
        "text": 50,
        "translation": 30,
        "suggested_seconds": 11,
        "exam_kind": 18,
        "exam_level": 14,
        "cue_card_bullets": 38,
        "prep_seconds": 11,
        "audio_url": 30,
    },
    TYPE_INSTRUCTION: {
        "title": 20,
        "text": 60,
        "suggested_seconds": 11,
    },
}


def build_template(fmt: str) -> tuple[bytes, str, str]:
    """生成导入模板：xlsx（分工作表 + 下拉）或 CSV（题型列 + 并集列）。"""
    if fmt == "csv":
        return _build_csv_template()
    return _build_xlsx_template()


def _build_xlsx_template() -> tuple[bytes, str, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)

    guide = workbook.create_sheet("说明")
    guide.column_dimensions["A"].width = 120
    for row_idx, line in enumerate(GUIDE_LINES, start=1):
        cell = guide.cell(row=row_idx, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for qtype in QUESTION_BANK_TYPE_ORDER:
        columns = COLUMNS_BY_TYPE[qtype]
        sheet = workbook.create_sheet(TYPE_LABELS[qtype])
        for col_idx, spec in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=spec.header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(col_idx)].width = TEMPLATE_WIDTHS[
                qtype
            ][spec.field]
        sheet.freeze_panes = "A2"
        for row_idx, example in enumerate(TEMPLATE_EXAMPLES[qtype], start=2):
            for col_idx, value in enumerate(example, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                cell.font = Font(italic=True, color="FF808080")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for fieldname, options in TEMPLATE_DROPDOWNS[qtype].items():
            col_idx = next(
                i for i, spec in enumerate(columns, start=1) if spec.field == fieldname
            )
            validation = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
            )
            validation.error = "请从下拉列表选择 Please pick from the list"
            validation.errorTitle = "无效选项 Invalid option"
            sheet.add_data_validation(validation)
            letter = get_column_letter(col_idx)
            validation.add(f"{letter}2:{letter}1001")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "question-bank-import-template.xlsx", XLSX_MEDIA_TYPE


CSV_TEMPLATE_EXAMPLES: list[dict[str, str]] = [
    {
        "__type__": "文章朗读",
        "title": "My Morning Routine",
        "unit_title": "Unit 1",
        "topic": "Daily Life",
        "cefr_band": "B1",
        "text": "I usually wake up at six o'clock. After breakfast, I walk to school with my neighbour.",
        "translation": "我通常六点起床。早饭后我和邻居一起走路去上学。",
        "suggested_seconds": "45",
        "is_active": "启用",
    },
    {
        "__type__": "听句复述",
        "text": "The library opens at nine o'clock.",
        "translation": "图书馆九点开门。",
        "suggested_seconds": "8",
        "replay_limit": "3",
    },
    {
        "__type__": "情景问答",
        "topic": "Travel",
        "text": "Do you prefer travelling alone or with friends? Why?",
        "translation": "你喜欢独自旅行还是和朋友一起？为什么？",
        "suggested_seconds": "20",
    },
    {
        "__type__": "题目说明",
        "title": "Part B · 开始",
        "text": "In this part, you will talk about the topic on the card. You will have one minute to prepare.",
        "suggested_seconds": "60",
    },
]


def _build_csv_template() -> tuple[bytes, str, str]:
    header = [CSV_TYPE_COLUMN.zh] + [c.zh for c in CSV_UNION_COLUMNS]
    rows = [header]
    for example in CSV_TEMPLATE_EXAMPLES:
        rows.append(
            [
                example.get("__type__", ""),
                *[example.get(c.field, "") for c in CSV_UNION_COLUMNS],
            ]
        )
    return (
        _csv_bytes(rows),
        "question-bank-import-template.csv",
        CSV_MEDIA_TYPE,
    )


def _csv_guard(value: str) -> str:
    """CSV 公式注入防护：与前端 lib/csv.ts 同规则（= +/-/@/制表符/回车加 ' 前缀）。"""
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def _csv_unguard(cell: str) -> str:
    """_csv_guard 的逆运算：还原导出时加的公式防护前缀。

    规则与 guard 严格镜像——仅当首字符是 `'` 且第二字符属于危险字符集时
    才剥掉前缀（内容本就以 `'=` 开头的概率极低，且 Excel 中该写法本来就
    表示转义）。没有这一步，导出→导入会永久改变以 - / = / + 开头的
    题干并凭空新增重复题（评审 P2："- Explain your choice." → "'- …"）。
    """
    if len(cell) >= 2 and cell[0] == "'" and cell[1] in "=+-@\t\r":
        return cell[1:]
    return cell


def _csv_bytes(rows: list[list[str]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows([[_csv_guard(c) for c in row] for row in rows])
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


# ── 批量导出（筛选条件镜像现有列表接口；格式与导入模板兼容） ────────


def export_bank(
    session: Session,
    *,
    types: list[str] | None,
    topic: str | None,
    q: str | None,
    band: str | None,
    exam_kind: str | None,
    exam_level: str | None,
    active: bool | None,
    fmt: str,
) -> tuple[bytes, str, str]:
    selected = list(types) if types else list(QUESTION_BANK_TYPE_ORDER)
    unknown = [t for t in selected if t not in TYPE_LABELS]
    if unknown:
        raise HTTPException(status_code=422, detail=f"未知题型：{'/'.join(unknown)}")
    band_upper = band.strip().upper() if band else None
    if band_upper is not None and band_upper not in VALID_BANDS:
        raise HTTPException(status_code=422, detail="难度无效，可选：A2/B1/B2")

    rows_by_type: dict[str, list[dict[str, str]]] = {}
    if TYPE_PASSAGE in selected:
        rows_by_type[TYPE_PASSAGE] = _export_passages(
            session, topic, q, band_upper, active
        )
    if TYPE_REPEAT in selected:
        rows_by_type[TYPE_REPEAT] = _export_sentences(session, exam_kind, exam_level)
    if TYPE_QUESTION in selected:
        rows_by_type[TYPE_QUESTION] = _export_questions(
            session, topic, q, band_upper, exam_kind, exam_level
        )
    if TYPE_INSTRUCTION in selected:
        rows_by_type[TYPE_INSTRUCTION] = _export_instructions(session)

    total = sum(len(v) for v in rows_by_type.values())
    if total > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"导出行数超过上限（{MAX_EXPORT_ROWS} 行），请按题型/主题等筛选分批导出"
                "（导出文件保持可直接回导）"
            ),
        )

    date_tag = get_datetime_utc().date().isoformat()
    if fmt == "csv":
        # CSV 行必须按 CSV_UNION_COLUMNS 的字段名取值对齐——各题型自己的
        # 列顺序与并集列顺序不同，直接拼接会整体错位（评审 P1：回导损坏题目）
        header = [CSV_TYPE_COLUMN.zh] + [c.zh for c in CSV_UNION_COLUMNS]
        rows = [header]
        for qtype in QUESTION_BANK_TYPE_ORDER:
            for row in rows_by_type.get(qtype, []):
                rows.append(
                    [
                        TYPE_LABELS[qtype],
                        *[row.get(c.field, "") for c in CSV_UNION_COLUMNS],
                    ]
                )
        return (
            _csv_bytes(rows),
            f"question-bank-export-{date_tag}.csv",
            CSV_MEDIA_TYPE,
        )
    return (
        _export_xlsx(rows_by_type),
        f"question-bank-export-{date_tag}.xlsx",
        XLSX_MEDIA_TYPE,
    )


def _export_passages(
    session: Session,
    topic: str | None,
    q: str | None,
    band_upper: str | None,
    active: bool | None,
) -> list[dict[str, str]]:
    conditions: list[Any] = [col(Passage.parent_passage_id).is_(None)]
    if q and q.strip():
        needle = q.strip()
        conditions.append(
            col(Passage.title).icontains(needle) | col(Passage.topic).icontains(needle)  # type: ignore[operator]
        )
    if topic is not None and topic.strip():
        conditions.append(Passage.topic == topic.strip())
    if band_upper is not None:
        conditions.append(Passage.cefr_band == band_upper)
    if active is not None:
        conditions.append(Passage.is_active == active)
    passages = session.exec(
        select(Passage)
        .where(*conditions)
        .order_by(col(Passage.created_at), col(Passage.id))
    ).all()
    unit_titles = {u.id: u.title for u in session.exec(select(Unit)).all()}
    return [
        {
            "title": p.title,
            "unit_title": unit_titles.get(p.unit_id, "") if p.unit_id else "",
            "topic": p.topic or "",
            "cefr_band": p.cefr_band,
            "text": p.text,
            "translation": p.translation or "",
            "suggested_seconds": str(p.suggested_seconds),
            "audio_url": p.audio_url or "",
            "is_active": "启用" if p.is_active else "停用",
        }
        for p in passages
    ]


def _export_sentences(
    session: Session, exam_kind: str | None, exam_level: str | None
) -> list[dict[str, str]]:
    stmt = select(RepeatSentence)
    if exam_kind is not None:
        stmt = stmt.where(RepeatSentence.exam_kind == exam_kind)  # type: ignore[arg-type]
    if exam_level is not None:
        stmt = stmt.where(RepeatSentence.exam_level == exam_level)  # type: ignore[arg-type]
    sentences = session.exec(
        stmt.order_by(col(RepeatSentence.created_at), col(RepeatSentence.id))
    ).all()
    title_map = dict(session.exec(select(Passage.id, Passage.title)).all())
    return [
        {
            "text": s.text,
            "translation": s.translation or "",
            "passage_title": title_map.get(s.passage_id, "") if s.passage_id else "",
            "suggested_seconds": str(s.suggested_seconds),
            "replay_limit": str(s.replay_limit),
            "exam_kind": s.exam_kind or "",
            "exam_level": s.exam_level or "",
            "audio_url": s.audio_url or "",
        }
        for s in sentences
    ]


def _export_questions(
    session: Session,
    topic: str | None,
    q: str | None,
    band_upper: str | None,
    exam_kind: str | None,
    exam_level: str | None,
) -> list[dict[str, str]]:
    conditions: list[Any] = []
    if topic is not None and topic.strip():
        conditions.append(Scenario.topic == topic.strip())
    if band_upper is not None:
        conditions.append(ScenarioQuestion.band == band_upper)
    if q and q.strip():
        needle = q.strip()
        conditions.append(
            col(ScenarioQuestion.text).icontains(needle)  # type: ignore[operator]
            | col(ScenarioQuestion.translation).icontains(needle)
        )
    if exam_kind is not None:
        conditions.append(ScenarioQuestion.exam_kind == exam_kind)
    if exam_level is not None:
        conditions.append(ScenarioQuestion.exam_level == exam_level)
    rows = session.exec(
        select(ScenarioQuestion, Scenario.topic)
        .join(Scenario, ScenarioQuestion.scenario_id == Scenario.id)  # ty: ignore[invalid-argument-type]
        .where(*conditions)
        .order_by(
            col(Scenario.topic),
            col(ScenarioQuestion.order_index),
            col(ScenarioQuestion.id),
        )
    ).all()
    return [
        {
            "topic": topic_name,
            "text": question.text,
            "translation": question.translation or "",
            "suggested_seconds": str(question.suggested_seconds),
            "exam_kind": question.exam_kind or "",
            "exam_level": question.exam_level or "",
            "cue_card_bullets": "\n".join(question.cue_card_bullets or ()),
            "prep_seconds": (
                str(question.prep_seconds) if question.prep_seconds is not None else ""
            ),
            "audio_url": question.audio_url or "",
        }
        for question, topic_name in rows
    ]


def _export_instructions(session: Session) -> list[dict[str, str]]:
    instructions = session.exec(
        select(Instruction).order_by(col(Instruction.created_at), col(Instruction.id))
    ).all()
    return [
        {
            "title": i.title or "",
            "text": i.text,
            "suggested_seconds": str(i.suggested_seconds),
        }
        for i in instructions
    ]


def _export_xlsx(rows_by_type: dict[str, list[dict[str, str]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    # Excel 单元格字符串硬上限 32767，openpyxl 写超长会**静默截断**——
    # 导出必须无损，先整体校验再生成（评审 P2：截断后仍返回成功）
    for qtype in QUESTION_BANK_TYPE_ORDER:
        for row in rows_by_type.get(qtype, []):
            for field_key, value in row.items():
                if len(value) > EXCEL_CELL_MAX_CHARS:
                    label = TYPE_LABELS[qtype]
                    raise HTTPException(
                        status_code=422,
                        detail=(
                            f"{label}「{row.get('title') or row.get('text', '')[:30]}」"
                            f"的{field_key}超过 Excel 单元格 {EXCEL_CELL_MAX_CHARS} 字符上限，"
                            "请改用 CSV 格式导出（无此限制，内容无损）"
                        ),
                    )

    workbook = Workbook()
    workbook.remove(workbook.active)
    for qtype in QUESTION_BANK_TYPE_ORDER:
        if qtype not in rows_by_type:
            continue
        columns = COLUMNS_BY_TYPE[qtype]
        sheet = workbook.create_sheet(TYPE_LABELS[qtype])
        for col_idx, spec in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=spec.header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(col_idx)].width = TEMPLATE_WIDTHS[
                qtype
            ][spec.field]
        sheet.freeze_panes = "A2"
        for row_idx, row in enumerate(rows_by_type[qtype], start=2):
            for col_idx, spec in enumerate(columns, start=1):
                value = row.get(spec.field, "")
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                # openpyxl 会把 "=" 开头的字符串当公式写入；导出内容必须恒为文本
                if value.startswith("="):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
