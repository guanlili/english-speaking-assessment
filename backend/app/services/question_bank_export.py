"""题库导出/模板/错误行文件（从 question_bank_io 拆分）。

与导入管线共享的常量、列规格与 csv 工具仍在 question_bank_io——
本模块单向依赖它（build_error_file 复用 plan_import 重建计划），避免循环。
"""

import io
from typing import Any

from fastapi import HTTPException
from sqlmodel import Session, col, select

from app.models import (
    VOCAB_LEVEL_ORDER,
    Instruction,
    Passage,
    RepeatSentence,
    Scenario,
    ScenarioQuestion,
    Unit,
    User,
    get_datetime_utc,
)
from app.services.question_bank_io import (
    COLUMNS_BY_TYPE,
    CSV_MEDIA_TYPE,
    CSV_TYPE_COLUMN,
    CSV_UNION_COLUMNS,
    EXCEL_CELL_MAX_CHARS,
    MAX_EXPORT_ROWS,
    MAX_UPLOAD_BYTES,
    QUESTION_BANK_TYPE_ORDER,
    TYPE_INSTRUCTION,
    TYPE_LABELS,
    TYPE_PASSAGE,
    TYPE_QUESTION,
    TYPE_REPEAT,
    VALID_BANDS,
    XLSX_MEDIA_TYPE,
    ErrorRow,
    ParsedFile,
    _csv_bytes,
    plan_import,
)

# ── 模板生成 ─────────────────────────────────────────────────────────

GUIDE_LINES = (
    "题库批量导入模板 · Question Bank Import Template",
    "",
    "使用步骤 Steps：",
    "1. 按题型填写对应工作表（每个工作表只放一种题型）Fill one sheet per question type",
    "2. 上传到「题目库 → 批量导入」预览 Upload in Question Bank → Bulk import and preview",
    "3. 确认无误后点「确认导入」 Confirm to import",
    "",
    "必填列 Required columns（其余列留空即用默认值；灰色示例行可删除）Leave other columns blank for defaults; example rows can be deleted：",
    "  文章朗读 Read Aloud：标题、正文 Title, Text",
    "  听句复述 Listen & Repeat：英文句子 Sentence",
    "  情景问答 Scenario Q&A：主题、英文题目 Topic, Question",
    "  题目说明 Instruction：说明文字 Text",
    "",
    "下拉选项 Dropdowns：难度/考试题型/考试级别/可听次数/状态请从下拉列表选择",
    "CSV 导入：单文件用「题型」列区分题型（值同工作表名）For CSV, use the Type column",
    "",
    "重复处理 Duplicates：完全重复自动跳过；疑似重复（内容相同、配置不同）导入时会请你确认",
    "Exact duplicates are skipped automatically; near duplicates need your confirmation",
    "导入不会修改已有题目 Import never modifies existing questions",
    "",
    "音频 Audio：先导入文字，再在题目编辑里上传音频文件；「音频链接」列只填本站链接",
    "Import text first, then upload audio files in the question editor; the Audio column only accepts this site's links",
    "",
    "上限 Limits：导入 2000 行 / 5MB；导出同受 2000 行与 5MB 限制（超出请按题型/主题分批导出）",
    "Limits: import 2000 rows / 5MB; exports share the same 2000-row and 5MB caps (split by type/topic when larger)",
    "编码 Encoding：Excel(.xlsx) 直接上传；CSV 需 UTF-8（推荐用本站下载的 CSV 模板）",
    "",
    "注意：本模板用于题目内容交换，不是完整备份——拆句朗读、逐题计时、题目说明的考试音频等高级配置不入模板，导入后如需请在题目编辑页设置",
    "Note: this template exchanges question content, not a full backup — advanced settings (per-sentence split, per-item timing, exam audio on instructions) must be reconfigured in the editor after import",
)

TEMPLATE_EXAMPLES: dict[str, list[list[str]]] = {
    TYPE_PASSAGE: [
        [
            "My Morning Routine",
            "Unit 1",
            "Daily Life",
            "B1",
            "I usually wake up at six o'clock. After breakfast, I walk to school with my neighbour.",
            "我通常六点起床。早饭后我和邻居一起走路去上学。",
            "45",
            "",
            "启用",
        ],
    ],
    TYPE_REPEAT: [
        [
            "The library opens at nine o'clock.",
            "图书馆九点开门。",
            "",
            "8",
            "3",
            "",
            "",
            "",
        ],
    ],
    TYPE_QUESTION: [
        [
            "Travel",
            "Do you prefer travelling alone or with friends? Why?",
            "你喜欢独自旅行还是和朋友一起？为什么？",
            "20",
            "",
            "",
            "",
            "",
            "",
        ],
        [
            "Travel",
            "Describe a place you visited that you would like to visit again.",
            "描述一个你去过、还想再去的地方。",
            "120",
            "IELTS · Part 2",
            "PET",
            "what the place is\nwhen you went there\nwhy you like it",
            "60",
            "",
        ],
    ],
    TYPE_INSTRUCTION: [
        [
            "Part B · 开始",
            "In this part, you will talk about the topic on the card. You will have one minute to prepare.",
            "60",
        ],
    ],
}

TEMPLATE_DROPDOWNS: dict[str, dict[str, list[str]]] = {
    TYPE_PASSAGE: {
        "cefr_band": ["A2", "B1", "B2"],
        "is_active": ["启用", "停用"],
    },
    TYPE_REPEAT: {
        "replay_limit": [str(i) for i in range(10)],
        "exam_kind": ["TOEFL · 听令复述"],
        "exam_level": list(VOCAB_LEVEL_ORDER),
    },
    TYPE_QUESTION: {
        "exam_kind": [
            "TOEFL · 情景访谈",
            "IELTS · Part 1",
            "IELTS · Part 2",
            "IELTS · Part 3",
        ],
        "exam_level": list(VOCAB_LEVEL_ORDER),
    },
    TYPE_INSTRUCTION: {},
}

TEMPLATE_WIDTHS: dict[str, dict[str, int]] = {
    TYPE_PASSAGE: {
        "title": 26,
        "unit_title": 14,
        "topic": 16,
        "cefr_band": 9,
        "text": 60,
        "translation": 34,
        "suggested_seconds": 11,
        "audio_url": 30,
        "is_active": 9,
    },
    TYPE_REPEAT: {
        "text": 50,
        "translation": 30,
        "passage_title": 20,
        "suggested_seconds": 11,
        "replay_limit": 11,
        "exam_kind": 18,
        "exam_level": 14,
        "audio_url": 30,
    },
    TYPE_QUESTION: {
        "topic": 16,
        "text": 50,
        "translation": 30,
        "suggested_seconds": 11,
        "exam_kind": 18,
        "exam_level": 14,
        "cue_card_bullets": 38,
        "prep_seconds": 11,
        "audio_url": 30,
    },
    TYPE_INSTRUCTION: {
        "title": 20,
        "text": 60,
        "suggested_seconds": 11,
    },
}


def build_template(fmt: str) -> tuple[bytes, str, str]:
    """生成导入模板：xlsx（分工作表 + 下拉）或 CSV（题型列 + 并集列）。"""
    if fmt == "csv":
        return _build_csv_template()
    return _build_xlsx_template()


def _build_xlsx_template() -> tuple[bytes, str, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    workbook.remove(workbook.active)

    guide = workbook.create_sheet("说明")
    guide.column_dimensions["A"].width = 120
    for row_idx, line in enumerate(GUIDE_LINES, start=1):
        cell = guide.cell(row=row_idx, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for qtype in QUESTION_BANK_TYPE_ORDER:
        columns = COLUMNS_BY_TYPE[qtype]
        sheet = workbook.create_sheet(TYPE_LABELS[qtype])
        for col_idx, spec in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=spec.header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(col_idx)].width = TEMPLATE_WIDTHS[
                qtype
            ][spec.field]
        sheet.freeze_panes = "A2"
        for row_idx, example in enumerate(TEMPLATE_EXAMPLES[qtype], start=2):
            for col_idx, value in enumerate(example, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                cell.font = Font(italic=True, color="FF808080")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for fieldname, options in TEMPLATE_DROPDOWNS[qtype].items():
            col_idx = next(
                i for i, spec in enumerate(columns, start=1) if spec.field == fieldname
            )
            validation = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
            )
            validation.error = "请从下拉列表选择 Please pick from the list"
            validation.errorTitle = "无效选项 Invalid option"
            sheet.add_data_validation(validation)
            letter = get_column_letter(col_idx)
            validation.add(f"{letter}2:{letter}1001")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "question-bank-import-template.xlsx", XLSX_MEDIA_TYPE


CSV_TEMPLATE_EXAMPLES: list[dict[str, str]] = [
    {
        "__type__": "文章朗读",
        "title": "My Morning Routine",
        "unit_title": "Unit 1",
        "topic": "Daily Life",
        "cefr_band": "B1",
        "text": "I usually wake up at six o'clock. After breakfast, I walk to school with my neighbour.",
        "translation": "我通常六点起床。早饭后我和邻居一起走路去上学。",
        "suggested_seconds": "45",
        "is_active": "启用",
    },
    {
        "__type__": "听句复述",
        "text": "The library opens at nine o'clock.",
        "translation": "图书馆九点开门。",
        "suggested_seconds": "8",
        "replay_limit": "3",
    },
    {
        "__type__": "情景问答",
        "topic": "Travel",
        "text": "Do you prefer travelling alone or with friends? Why?",
        "translation": "你喜欢独自旅行还是和朋友一起？为什么？",
        "suggested_seconds": "20",
    },
    {
        "__type__": "题目说明",
        "title": "Part B · 开始",
        "text": "In this part, you will talk about the topic on the card. You will have one minute to prepare.",
        "suggested_seconds": "60",
    },
]


def _build_csv_template() -> tuple[bytes, str, str]:
    header = [CSV_TYPE_COLUMN.zh] + [c.zh for c in CSV_UNION_COLUMNS]
    rows = [header]
    for example in CSV_TEMPLATE_EXAMPLES:
        rows.append(
            [
                example.get("__type__", ""),
                *[example.get(c.field, "") for c in CSV_UNION_COLUMNS],
            ]
        )
    return (
        _csv_bytes(rows),
        "question-bank-import-template.csv",
        CSV_MEDIA_TYPE,
    )


# 危险字符集（公式注入）+ 引号本身：原文以 ' 开头也要转义（见 _csv_guard）


# ── 待修正数据下载（错误行原始字段 + 辅助列；修正后可直接回传） ─────


ERROR_FILE_LINE_HEADER = "原行号\nOriginal line"
ERROR_FILE_REASON_HEADER = "错误原因\nError reason"


def build_error_file(
    session: Session, parsed: ParsedFile, editor: User | None = None
) -> tuple[bytes, str, str]:
    """把当前文件的全部错误行（解析期 + 校验期）生成可回传的修正文件。

    - 与上传格式同构（xlsx→xlsx / CSV→CSV），列用标准模板表头 + 两列辅助
      （原行号/错误原因）——辅助列表头不在解析别名表内，回传时被忽略；
    - 完整清单，不受预览界面 200 条截断限制；
    - 单元格值走同一套 guard/转义规则，原文引号、换行不损坏。
    """
    plan = plan_import(session, parsed, editor=editor)
    error_rows = plan.error_rows
    if not error_rows:
        raise HTTPException(status_code=422, detail="没有需要修正的行")
    if parsed.fmt == "csv":
        rows = [
            [
                CSV_TYPE_COLUMN.zh,
                *[c.zh for c in CSV_UNION_COLUMNS],
                "原行号",
                "错误原因",
            ]
        ]
        for error in error_rows:
            type_label = error.type_cell or TYPE_LABELS.get(error.qtype, "")
            rows.append(
                [
                    type_label,
                    *[error.cells.get(c.field, "") for c in CSV_UNION_COLUMNS],
                    str(error.line),
                    error.reason,
                ]
            )
        return (
            _csv_bytes(rows),
            "question-bank-fix.csv",
            CSV_MEDIA_TYPE,
        )
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    by_type: dict[str, list[ErrorRow]] = {}
    for error in error_rows:
        if error.qtype == "-":
            # xlsx 上传不会有题型列无效行（那类错误在文件级 422）；防御归并
            by_type.setdefault(TYPE_INSTRUCTION, []).append(error)
        else:
            by_type.setdefault(error.qtype, []).append(error)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for qtype in QUESTION_BANK_TYPE_ORDER:
        errors = by_type.get(qtype)
        if not errors:
            continue
        columns = COLUMNS_BY_TYPE[qtype]
        sheet = workbook.create_sheet(TYPE_LABELS[qtype])
        headers = [c.header for c in columns] + [
            ERROR_FILE_LINE_HEADER,
            ERROR_FILE_REASON_HEADER,
        ]
        for col_idx, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            width = TEMPLATE_WIDTHS[qtype].get(
                columns[col_idx - 1].field if col_idx <= len(columns) else "", 12
            )
            sheet.column_dimensions[get_column_letter(col_idx)].width = (
                14 if col_idx > len(columns) else width
            )
        sheet.freeze_panes = "A2"
        for row_idx, error in enumerate(errors, start=2):
            values = [error.cells.get(c.field, "") for c in columns] + [
                str(error.line),
                error.reason,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                if value.startswith("="):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "question-bank-fix.xlsx", XLSX_MEDIA_TYPE


# ── 批量导出（筛选条件镜像现有列表接口；格式与导入模板兼容） ────────


def export_bank(
    session: Session,
    *,
    types: list[str] | None,
    topic: str | None,
    q: str | None,
    band: str | None,
    exam_kind: str | None,
    exam_level: str | None,
    active: bool | None,
    fmt: str,
) -> tuple[bytes, str, str]:
    selected = list(types) if types else list(QUESTION_BANK_TYPE_ORDER)
    unknown = [t for t in selected if t not in TYPE_LABELS]
    if unknown:
        raise HTTPException(status_code=422, detail=f"未知题型：{'/'.join(unknown)}")
    band_upper = band.strip().upper() if band else None
    if band_upper is not None and band_upper not in VALID_BANDS:
        raise HTTPException(status_code=422, detail="难度无效，可选：A2/B1/B2")

    rows_by_type: dict[str, list[dict[str, str]]] = {}
    if TYPE_PASSAGE in selected:
        rows_by_type[TYPE_PASSAGE] = _export_passages(
            session, topic, q, band_upper, active
        )
    if TYPE_REPEAT in selected:
        rows_by_type[TYPE_REPEAT] = _export_sentences(session, exam_kind, exam_level)
    if TYPE_QUESTION in selected:
        rows_by_type[TYPE_QUESTION] = _export_questions(
            session, topic, q, band_upper, exam_kind, exam_level
        )
    if TYPE_INSTRUCTION in selected:
        rows_by_type[TYPE_INSTRUCTION] = _export_instructions(session)

    total = sum(len(v) for v in rows_by_type.values())
    if total > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"导出行数超过上限（{MAX_EXPORT_ROWS} 行），请按题型/主题等筛选分批导出"
                "（导出文件保持可直接回导）"
            ),
        )

    date_tag = get_datetime_utc().date().isoformat()
    if fmt == "csv":
        # CSV 行必须按 CSV_UNION_COLUMNS 的字段名取值对齐——各题型自己的
        # 列顺序与并集列顺序不同，直接拼接会整体错位（评审 P1：回导损坏题目）
        header = [CSV_TYPE_COLUMN.zh] + [c.zh for c in CSV_UNION_COLUMNS]
        rows = [header]
        for qtype in QUESTION_BANK_TYPE_ORDER:
            for row in rows_by_type.get(qtype, []):
                rows.append(
                    [
                        TYPE_LABELS[qtype],
                        *[row.get(c.field, "") for c in CSV_UNION_COLUMNS],
                    ]
                )
        data, filename, media_type = (
            _csv_bytes(rows),
            f"question-bank-export-{date_tag}.csv",
            CSV_MEDIA_TYPE,
        )
    else:
        data, filename, media_type = (
            _export_xlsx(rows_by_type),
            f"question-bank-export-{date_tag}.xlsx",
            XLSX_MEDIA_TYPE,
        )
    # 行数对齐之外还要对齐体积：导入有 5MB 上限，超体积的导出文件回导会
    # 被 413 挡下（评审 P2：1800 篇 ×3KB 导出 5.9MB）。超限直接 422 提示
    # 分批，「导出文件可直接回导」的承诺才成立
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=422,
            detail=(
                f"导出文件约 {len(data) // 1024 // 1024}MB，超过 "
                f"{MAX_UPLOAD_BYTES // 1024 // 1024}MB 上传上限，"
                "请按题型/主题缩小范围分批导出（导出文件保持可直接回导）"
            ),
        )
    return data, filename, media_type


def _export_passages(
    session: Session,
    topic: str | None,
    q: str | None,
    band_upper: str | None,
    active: bool | None,
) -> list[dict[str, str]]:
    conditions: list[Any] = [col(Passage.parent_passage_id).is_(None)]
    if q and q.strip():
        needle = q.strip()
        conditions.append(
            col(Passage.title).icontains(needle) | col(Passage.topic).icontains(needle)  # type: ignore[operator]
        )
    if topic is not None and topic.strip():
        conditions.append(Passage.topic == topic.strip())
    if band_upper is not None:
        conditions.append(Passage.cefr_band == band_upper)
    if active is not None:
        conditions.append(Passage.is_active == active)
    # MAX_EXPORT_ROWS+1 探测：超限时先于全量组装触发 422，且截断到
    # MAX+1 就足以精确判定超限（任一题型真行数 > MAX 时截断值必 > MAX）
    passages = session.exec(
        select(Passage)
        .where(*conditions)
        .order_by(col(Passage.created_at), col(Passage.id))
        .limit(MAX_EXPORT_ROWS + 1)
    ).all()
    unit_ids = {p.unit_id for p in passages if p.unit_id is not None}
    unit_titles = (
        {
            u.id: u.title
            for u in session.exec(select(Unit).where(col(Unit.id).in_(unit_ids))).all()
        }
        if unit_ids
        else {}
    )
    return [
        {
            "title": p.title,
            "unit_title": unit_titles.get(p.unit_id, "") if p.unit_id else "",
            "topic": p.topic or "",
            "cefr_band": p.cefr_band,
            "text": p.text,
            "translation": p.translation or "",
            "suggested_seconds": str(p.suggested_seconds),
            "audio_url": p.audio_url or "",
            "is_active": "启用" if p.is_active else "停用",
        }
        for p in passages
    ]


def _export_sentences(
    session: Session, exam_kind: str | None, exam_level: str | None
) -> list[dict[str, str]]:
    stmt = select(RepeatSentence)
    if exam_kind is not None:
        stmt = stmt.where(RepeatSentence.exam_kind == exam_kind)  # type: ignore[arg-type]
    if exam_level is not None:
        stmt = stmt.where(RepeatSentence.exam_level == exam_level)  # type: ignore[arg-type]
    sentences = session.exec(
        stmt.order_by(col(RepeatSentence.created_at), col(RepeatSentence.id)).limit(
            MAX_EXPORT_ROWS + 1
        )
    ).all()
    passage_ids = {s.passage_id for s in sentences if s.passage_id is not None}
    title_map = (
        dict(
            session.exec(
                select(Passage.id, Passage.title).where(
                    col(Passage.id).in_(passage_ids)
                )
            ).all()
        )
        if passage_ids
        else {}
    )
    return [
        {
            "text": s.text,
            "translation": s.translation or "",
            "passage_title": title_map.get(s.passage_id, "") if s.passage_id else "",
            "suggested_seconds": str(s.suggested_seconds),
            "replay_limit": str(s.replay_limit),
            "exam_kind": s.exam_kind or "",
            "exam_level": s.exam_level or "",
            "audio_url": s.audio_url or "",
        }
        for s in sentences
    ]


def _export_questions(
    session: Session,
    topic: str | None,
    q: str | None,
    band_upper: str | None,
    exam_kind: str | None,
    exam_level: str | None,
) -> list[dict[str, str]]:
    conditions: list[Any] = []
    if topic is not None and topic.strip():
        conditions.append(Scenario.topic == topic.strip())
    if band_upper is not None:
        conditions.append(ScenarioQuestion.band == band_upper)
    if q and q.strip():
        needle = q.strip()
        conditions.append(
            col(ScenarioQuestion.text).icontains(needle)  # type: ignore[operator]
            | col(ScenarioQuestion.translation).icontains(needle)
        )
    if exam_kind is not None:
        conditions.append(ScenarioQuestion.exam_kind == exam_kind)
    if exam_level is not None:
        conditions.append(ScenarioQuestion.exam_level == exam_level)
    rows = session.exec(
        select(ScenarioQuestion, Scenario.topic)
        .join(Scenario, ScenarioQuestion.scenario_id == Scenario.id)  # ty: ignore[invalid-argument-type]
        .where(*conditions)
        .order_by(
            col(Scenario.topic),
            col(ScenarioQuestion.order_index),
            col(ScenarioQuestion.id),
        )
        .limit(MAX_EXPORT_ROWS + 1)
    ).all()
    return [
        {
            "topic": topic_name,
            "text": question.text,
            "translation": question.translation or "",
            "suggested_seconds": str(question.suggested_seconds),
            "exam_kind": question.exam_kind or "",
            "exam_level": question.exam_level or "",
            "cue_card_bullets": "\n".join(question.cue_card_bullets or ()),
            "prep_seconds": (
                str(question.prep_seconds) if question.prep_seconds is not None else ""
            ),
            "audio_url": question.audio_url or "",
        }
        for question, topic_name in rows
    ]


def _export_instructions(session: Session) -> list[dict[str, str]]:
    instructions = session.exec(
        select(Instruction)
        .order_by(col(Instruction.created_at), col(Instruction.id))
        .limit(MAX_EXPORT_ROWS + 1)
    ).all()
    return [
        {
            "title": i.title or "",
            "text": i.text,
            "suggested_seconds": str(i.suggested_seconds),
        }
        for i in instructions
    ]


def _export_xlsx(rows_by_type: dict[str, list[dict[str, str]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    # Excel 单元格字符串硬上限 32767，openpyxl 写超长会**静默截断**——
    # 导出必须无损，先整体校验再生成（评审 P2：截断后仍返回成功）
    for qtype in QUESTION_BANK_TYPE_ORDER:
        for row in rows_by_type.get(qtype, []):
            for field_key, value in row.items():
                if len(value) > EXCEL_CELL_MAX_CHARS:
                    label = TYPE_LABELS[qtype]
                    raise HTTPException(
                        status_code=422,
                        detail=(
                            f"{label}「{row.get('title') or row.get('text', '')[:30]}」"
                            f"的{field_key}超过 Excel 单元格 {EXCEL_CELL_MAX_CHARS} 字符上限，"
                            "请改用 CSV 格式导出（无此限制，内容无损）"
                        ),
                    )

    workbook = Workbook()
    workbook.remove(workbook.active)
    for qtype in QUESTION_BANK_TYPE_ORDER:
        if qtype not in rows_by_type:
            continue
        columns = COLUMNS_BY_TYPE[qtype]
        sheet = workbook.create_sheet(TYPE_LABELS[qtype])
        for col_idx, spec in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=spec.header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(col_idx)].width = TEMPLATE_WIDTHS[
                qtype
            ][spec.field]
        sheet.freeze_panes = "A2"
        for row_idx, row in enumerate(rows_by_type[qtype], start=2):
            for col_idx, spec in enumerate(columns, start=1):
                value = row.get(spec.field, "")
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                # openpyxl 会把 "=" 开头的字符串当公式写入；导出内容必须恒为文本
                if value.startswith("="):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
