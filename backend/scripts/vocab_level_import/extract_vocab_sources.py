"""五级词库资料盘点与清洗工具（本地运行，产出不入仓库）。

针对学校提供的 /Users/guanhongli/Downloads/2026Vocabulary(1)/ 原始资料：
KET / PET / 学术词汇 / 四级词汇 / 雅思&托福词汇（雅思词表、托福真词汇、
托福词组三个来源分别识别）。工具只做提取、规范化、去重与统计：

- 规范化词条写入 output/（已被 .gitignore 忽略），绝不提交
- 质量报告写入 docs/five-level-vocab-quality-report.md（只有统计、
  冲突清单与每来源 ≤5 行格式样例，不含完整词表）

版权约束：这些资料为出版物词表，未确认线上使用授权前不得导入生产库，
也不得把原始文件或提取出的完整词表提交到仓库。

用法（在 backend/ 目录）：
    uv run --with openpyxl --with xlrd python scripts/vocab_level_import/extract_vocab_sources.py \
        --source-dir "/Users/guanhongli/Downloads/2026Vocabulary(1)" --report

依赖：macOS 自带 textutil（.doc）、poppler 的 pdftoppm、tesseract（含 chi_sim）。
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

# 五级固定顺序（越靠前越容易）；实际难度默认取最早一级
LEVEL_ORDER = ["KET", "PET", "ACADEMIC", "CET4", "IELTS_TOEFL"]

DEFAULT_SOURCE_DIR = "/Users/guanhongli/Downloads/2026Vocabulary(1)"
OUTPUT_DIR = Path(__file__).parent / "output"
REPO_ROOT = Path(__file__).resolve().parents[3]
REPORT_PATH = REPO_ROOT / "docs" / "five-level-vocab-quality-report.md"

CJK_RE = re.compile(r"[\u4e00-\u9fff]")
VALID_HEADWORD_RE = re.compile(r"^[a-zA-Z][a-zA-Z'\-]*(?: [a-zA-Z][a-zA-Z'\-]*)*$")


def normalize_headword(raw: str) -> str:
    """与后端判分一致的规范化：NFKC + 首尾空白 + casefold。"""
    return unicodedata.normalize("NFKC", raw).strip().casefold()


@dataclass
class Row:
    """一条规范化词条。headword 是匹配键（规范化小写）。"""

    level: str
    headword: str  # 规范化（casefold）后的匹配形式
    part_of_speech: str | None = None
    meaning_zh: str | None = None
    is_phrase: bool = False
    needs_review: bool = False
    note: str | None = None
    sources: list[str] = field(default_factory=list)

    def merge(self, other: Row) -> None:
        """同 (level, headword, meaning) 合并：来源追加，空位补齐。"""
        for source in other.sources:
            if source not in self.sources:
                self.sources.append(source)
        if not self.part_of_speech and other.part_of_speech:
            self.part_of_speech = other.part_of_speech
        if not self.meaning_zh and other.meaning_zh:
            self.meaning_zh = other.meaning_zh
        self.needs_review = self.needs_review or other.needs_review
        if other.note and other.note not in (self.note or ""):
            self.note = f"{self.note}；{other.note}" if self.note else other.note


def _cells_to_list(ws, max_col: int = 8) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else str(c).strip() for c in row[:max_col]])
    return rows


def _parse_ket(path: Path) -> list[Row]:
    """KET 整理版 xlsx：总表（序号/天数/英语单词/词性/中文释义）。

    同词跨天重复 → 合并来源（source_label 带 DAY 标签）。
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["总表"]
    rows_out: list[Row] = []
    for cells in _cells_to_list(ws):
        if len(cells) < 5 or cells[0] in ("", "序号") or not cells[0].isdigit():
            continue
        day, word, pos, meaning = cells[1], cells[2], cells[3], cells[4]
        if not word:
            continue
        rows_out.append(
            Row(
                level="KET",
                headword=normalize_headword(word),
                part_of_speech=pos or None,
                meaning_zh=meaning or None,
                sources=[f"KET整理版·{day}" if day else "KET整理版"],
            )
        )
    wb.close()
    return rows_out


def _parse_pet_ocr(path: Path) -> list[Row]:
    """PET 扫描件：pdftoppm 转图 + tesseract OCR。

    英文词头与词性可靠；中文释义 OCR 不可靠 → 全部 needs_review，
    释义留空由人工核对补录。OCR 结果必须人工核对（设计要求）。
    """
    info = subprocess.run(
        ["pdfinfo", str(path)], capture_output=True, text=True, check=True
    ).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    rows_out: list[Row] = []
    entry_re = re.compile(r"^([a-zA-Z][a-zA-Z'\- ]*?)\s*\(([^)]*)\)\s*(.*)$")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["pdftoppm", "-r", "150", "-png", str(path), f"{tmp}/pet"],
            check=True,
            capture_output=True,
        )
        for image in sorted(Path(tmp).glob("pet-*.png")):
            text = subprocess.run(
                ["tesseract", str(image), "stdout", "-l", "eng"],
                capture_output=True,
                text=True,
                check=True,
            ).stdout
            for line in text.splitlines():
                line = line.strip()
                match = entry_re.match(line)
                if not match:
                    continue
                raw_headword, pos, rest = match.groups()
                # 「all right/alright」类拼法变体：取第一段为词头，其余入备注
                note = None
                headword = raw_headword.strip()
                if "/" in headword:
                    variants = [v.strip() for v in headword.split("/") if v.strip()]
                    headword = variants[0]
                    note = f"拼法变体（OCR）：{'/'.join(variants)}"
                if rest.strip():
                    note = (
                        f"{note}；OCR尾注：{rest.strip()}"
                        if note
                        else f"OCR尾注：{rest.strip()}"
                    )
                if not VALID_HEADWORD_RE.match(headword):
                    continue
                rows_out.append(
                    Row(
                        level="PET",
                        headword=normalize_headword(headword),
                        part_of_speech=pos.strip() or None,
                        needs_review=True,
                        note=note,
                        sources=["PET扫描件OCR"],
                    )
                )
    return rows_out


def _parse_academic(
    path: Path, sheet: str, source_label: str, skip_rows: int
) -> list[Row]:
    """学术词汇：两个 xlsx 是同一份表的两个整理版本（合并来源去重）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows_out: list[Row] = []
    for cells in _cells_to_list(ws):
        # 整理版：序号/书页/单词/音标/词性/常考释义；list 版：序号/单词/音标/词性/释义
        word_idx = 2 if "书页" in cells or len(cells) > 5 else 1
        if len(cells) <= word_idx or cells[0] in ("", "序号") or not cells[0].isdigit():
            continue
        word = cells[word_idx]
        pos = cells[word_idx + 2] if len(cells) > word_idx + 2 else ""
        meaning = cells[word_idx + 3] if len(cells) > word_idx + 3 else ""
        if not word:
            continue
        rows_out.append(
            Row(
                level="ACADEMIC",
                headword=normalize_headword(word),
                part_of_speech=pos or None,
                meaning_zh=meaning or None,
                sources=[source_label],
            )
        )
    wb.close()
    return rows_out


def _parse_cet4(path: Path) -> list[Row]:
    """四级：新东方乱序版 .doc（textutil 转 txt）。

    格式：Word List N 分节；词条行「单词 中文释义」（首个 CJK 字符处切分）。
    """
    text = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        text=True,
        check=True,
    ).stdout
    rows_out: list[Row] = []
    section = "四级乱序版"
    entry_re = re.compile(r"^([a-zA-Z][a-zA-Z'\- ]*?)\s+([\u4e00-\u9fff].*)$")
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        list_match = re.match(r"^Word List\s*(\d+)", line, re.IGNORECASE)
        if list_match:
            section = f"四级乱序版·List{list_match.group(1)}"
            continue
        match = entry_re.match(line)
        if not match:
            continue
        word, meaning = match.groups()
        if not VALID_HEADWORD_RE.match(word.strip()):
            continue
        rows_out.append(
            Row(
                level="CET4",
                headword=normalize_headword(word.strip()),
                meaning_zh=meaning.strip() or None,
                sources=[section],
            )
        )
    return rows_out


def _parse_ielts(path: Path) -> list[Row]:
    """雅思阅读词汇 xlsx：场景分节（节头行=单格含「类」，如 建筑类Architecture）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows_out: list[Row] = []
    scene = "未分节"
    for cells in _cells_to_list(ws, max_col=4):
        non_empty = [c for c in cells if c]
        if len(non_empty) == 1 and ("类" in non_empty[0] or "Scene" in non_empty[0]):
            scene = non_empty[0]
            continue
        if len(cells) >= 3 and cells[0] and cells[0] in ("英文单词", ""):
            continue
        if len(cells) >= 1 and cells[0] and VALID_HEADWORD_RE.match(cells[0].strip()):
            rows_out.append(
                Row(
                    level="IELTS_TOEFL",
                    headword=normalize_headword(cells[0].strip()),
                    part_of_speech=(cells[1] or None) if len(cells) > 1 else None,
                    meaning_zh=(cells[2] or None) if len(cells) > 2 else None,
                    sources=[f"雅思阅读词汇·{scene}"],
                )
            )
    wb.close()
    return rows_out


def _parse_toefl_words(path: Path) -> list[Row]:
    """托福真词汇 Chapter1-2：自带出版方 CEFR 标注（仅入备注，不作平台分级）。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["托福真词汇_完整词表"]
    rows_out: list[Row] = []
    for cells in _cells_to_list(ws):
        # 序号/章节/难度/单词/音标/词性/释义
        if len(cells) < 7 or cells[0] in ("", "序号") or not cells[0].isdigit():
            continue
        chapter, cefr, word = cells[1], cells[2], cells[3]
        pos, meaning = cells[5], cells[6] if len(cells) > 6 else ""
        if not word:
            continue
        rows_out.append(
            Row(
                level="IELTS_TOEFL",
                headword=normalize_headword(word),
                part_of_speech=pos or None,
                meaning_zh=meaning or None,
                note=f"出版方CEFR标注：{cefr}" if cefr else None,
                sources=[f"托福真词汇·{chapter}" if chapter else "托福真词汇"],
            )
        )
    wb.close()
    return rows_out


def _parse_toefl_phrases(path: Path) -> list[Row]:
    """猴哥托福词组 1200（.xls）：Sheet1，前几行是版本说明。词组单独识别。"""
    import xlrd

    book = xlrd.open_workbook(str(path))
    ws = book.sheet_by_index(0)
    rows_out: list[Row] = []
    for i in range(ws.nrows):
        cells = [str(ws.cell_value(i, j)).strip() for j in range(min(4, ws.ncols))]
        joined = " ".join(c for c in cells if c)
        # 版本说明行包含「版本/来源/更新/词组」等长文案，跳过
        if len(joined) > 60 or not any(cells):
            continue
        phrase = next((c for c in cells if c), "")
        if not phrase or " " not in phrase.strip():
            continue  # 词组表里混入的单词行：跳过（词组表只收词组）
        meaning = next((c for c in cells[1:] if CJK_RE.search(c)), "")
        rows_out.append(
            Row(
                level="IELTS_TOEFL",
                headword=normalize_headword(phrase),
                meaning_zh=meaning or None,
                is_phrase=True,
                sources=["猴哥托福词组1200"],
            )
        )
    return rows_out


def dedupe(rows: list[Row]) -> list[Row]:
    """批内去重：同 (level, headword, meaning) 合并；同词同级不同释义 = 同形异义，各留一行。"""
    merged: dict[tuple[str, str, str | None], Row] = {}
    order: list[tuple[str, str, str | None]] = []
    for row in rows:
        key = (row.level, row.headword, normalize_headword(row.meaning_zh or ""))
        if key in merged:
            merged[key].merge(row)
        else:
            merged[key] = row
            order.append(key)
    return [merged[key] for key in order]


def extract_all(source_dir: Path) -> dict[str, list[Row]]:
    """按来源提取全部词条（各自 dedupe），按五级固定顺序返回。"""
    by_level: dict[str, list[Row]] = {level: [] for level in LEVEL_ORDER}

    ket = dedupe(_parse_ket(source_dir / "1.KET单词表14天-整理版 副本.xlsx"))
    by_level["KET"] = ket

    by_level["PET"] = dedupe(_parse_pet_ocr(source_dir / "2. PET高频词汇表-新.pdf"))

    academic_a = dedupe(
        _parse_academic(
            source_dir / "3.基础学术词汇-整理版.xlsx",
            "词汇总表",
            "基础学术词汇·整理版",
            skip_rows=2,
        )
    )
    academic_b = dedupe(
        _parse_academic(
            source_dir / "3.academic_vocabulary_list.xlsx",
            "Academic Vocabulary",
            "academic_vocabulary_list",
            skip_rows=2,
        )
    )
    by_level["ACADEMIC"] = dedupe(academic_a + academic_b)

    by_level["CET4"] = dedupe(_parse_cet4(source_dir / "4.新东方四级词汇乱序版.doc"))
    by_level["IELTS_TOEFL"] = dedupe(
        _parse_ielts(source_dir / "5.《3 1雅思阅读词汇》（基础篇——场景核心词）.xlsx")
        + _parse_toefl_words(source_dir / "5.托福真词汇_Chapter1_Chapter2_词汇表.xlsx")
        + _parse_toefl_phrases(source_dir / "5.猴哥托福词组1200打印版7.0.xls")
    )
    return by_level


def cross_level_conflicts(by_level: dict[str, list[Row]]) -> list[dict]:
    """跨级冲突：同一 headword 出现在多个级别。实际难度将取最早（最易）一级。"""
    seen: dict[str, list[tuple[str, Row]]] = {}
    for level in LEVEL_ORDER:
        for row in by_level[level]:
            seen.setdefault(row.headword, []).append((level, row))
    conflicts = []
    for headword, hits in seen.items():
        if len({level for level, _ in hits}) > 1:
            conflicts.append(
                {
                    "headword": headword,
                    "levels": [level for level, _ in hits],
                    "effective_level": min(
                        (level for level, _ in hits), key=LEVEL_ORDER.index
                    ),
                    "meanings": [
                        {"level": level, "meaning": row.meaning_zh}
                        for level, row in hits
                    ][:6],
                }
            )
    return sorted(conflicts, key=lambda c: LEVEL_ORDER.index(c["effective_level"]))


def build_report(by_level: dict[str, list[Row]], conflicts: list[dict]) -> str:
    """质量报告：只有统计与样例（每来源 ≤5 行），不含完整词表。"""
    lines = [
        "# 五级词库 · 资料清洗与数据质量报告",
        "",
        "> 生成：`backend/scripts/vocab_level_import/extract_vocab_sources.py`（本地运行）。",
        "> 版权约束：原始教材与提取出的完整词表均不入仓库、未经授权不导入生产库；本报告只含统计与格式样例。",
        "",
        "## 五级定义（固定顺序，实际难度取最早一级）",
        "",
        "| 序 | 级别代码 | 名称 |",
        "|---|---|---|",
    ]
    labels = {
        "KET": "KET 词汇",
        "PET": "PET 词汇",
        "ACADEMIC": "学术词汇",
        "CET4": "四级词汇",
        "IELTS_TOEFL": "雅思&托福词汇",
    }
    for i, level in enumerate(LEVEL_ORDER, start=1):
        lines.append(f"| {i} | `{level}` | {labels[level]} |")

    lines += [
        "",
        "## 各级提取统计",
        "",
        "| 级别 | 去重后词条 | 词组 | 待人工核对 | 来源文件 |",
        "|---|---|---|---|---|",
    ]
    source_files = {
        "KET": "1.KET单词表14天-整理版 副本.xlsx（总表 623 词条口径，跨天重复词已合并）",
        "PET": "2. PET高频词汇表-新.pdf（42 页扫描件，OCR；xlsx 为空文件）",
        "ACADEMIC": "3.基础学术词汇-整理版.xlsx + 3.academic_vocabulary_list.xlsx（同一份表的重复来源，已合并）",
        "CET4": "4.新东方四级词汇乱序版.doc（textutil 提取）",
        "IELTS_TOEFL": "5.雅思阅读词汇.xlsx + 5.托福真词汇_Chapter1-2.xlsx + 5.猴哥托福词组1200.xls（三来源分别识别）",
    }
    for level in LEVEL_ORDER:
        rows = by_level[level]
        phrases = sum(1 for r in rows if r.is_phrase)
        review = sum(1 for r in rows if r.needs_review)
        lines.append(
            f"| {labels[level]} | {len(rows)} | {phrases} | {review} | {source_files[level]} |"
        )

    lines += [
        "",
        "## 同级重复与来源保留",
        "",
        "同 (级别, 单词, 释义) 的重复行已合并，来源标签保留在 `sources` 列表（如 KET 跨天重复：apartment 同时出现在 DAY9/DAY10 → sources 含两天标签）。同词同级不同释义 = 同形异义，各留一行（sense 区分）。",
        "",
    ]

    lines += [
        "## 跨级冲突（实际难度取最早一级）",
        "",
        f"共 **{len(conflicts)}** 个词出现在多个级别。前 30 个：",
        "",
        "| 单词 | 出现级别 | 实际难度 |",
        "|---|---|---|",
    ]
    for conflict in conflicts[:30]:
        lines.append(
            f"| {conflict['headword']} | {'、'.join(conflict['levels'])} | **{conflict['effective_level']}** |"
        )
    if len(conflicts) > 30:
        lines.append(
            f"| … | 共 {len(conflicts)} 个，完整清单见 output/cross-level-conflicts.json（本地） | |"
        )

    lines += ["", "## 待人工核对（needs_review）", ""]
    pet_rows = by_level["PET"]
    lines += [
        f"- **PET 全部 {len(pet_rows)} 条来自扫描件 OCR**：英文词头与词性可靠，中文释义 OCR 不可靠已留空，需人工补录核对后方可启用。",
        "- OCR 中识别出的拼法变体（如 all right/alright）已写入备注，不作为可接受拼写自动生效（同义词/变体不自动放宽是平台判分红线）。",
    ]

    lines += ["", "## 格式样例（每来源 ≤5 条，仅示意字段结构）", ""]
    sample_sources = {
        "KET": None,
        "PET": None,
        "ACADEMIC": None,
        "CET4": None,
        "IELTS_TOEFL": None,
    }
    for level in LEVEL_ORDER:
        lines.append(f"### {labels[level]}")
        lines.append("")
        rows = by_level[level][:5]
        lines.append("| headword | 词性 | 释义 | is_phrase | needs_review | sources |")
        lines.append("|---|---|---|---|---|---|")
        for row in rows:
            lines.append(
                f"| {row.headword} | {row.part_of_speech or '–'} | {row.meaning_zh or '–'} | {row.is_phrase} | {row.needs_review} | {'; '.join(row.sources)} |"
            )
        lines.append("")

    lines += [
        "## 尚需人工核对的行（汇总）",
        "",
        f"1. PET {len(pet_rows)} 条 OCR 行：释义留空待补，全部 `needs_review=true`；导入后可在管理端「五级词库」按该标记过滤逐条核对。",
        "2. 猴哥托福词组中混入的单词行已跳过；如发现词组解析缺行，对照原 xls 复核。",
        "3. 四级 .doc 为出版物全文提取（新东方《四级词汇词根+联想记忆法（乱序版）》），**未确认授权前不得导入生产库**。",
        "4. 学术词汇两份 xlsx 内容重叠度高（同一份表），合并后如数量异常请对照说明页。",
        "",
        "## 导入方式",
        "",
        "- 确认授权后：管理端「五级词库」逐级上传导入（预览→确认，事务性），或用本工具 output/ 的规范化 JSON 走 `apply_import.py`。",
        "- 导入只新增 `vocab_level_entry` 行；已发布任务快照与历史 A2/B1/B2 词汇分析结果不受影响。",
    ]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--report", action="store_true", help="生成质量报告到 docs/")
    args = parser.parse_args()

    source_dir = Path(args.source_dir)
    by_level = extract_all(source_dir)
    conflicts = cross_level_conflicts(by_level)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for level in LEVEL_ORDER:
        payload = [
            {
                "level": row.level,
                "headword": row.headword,
                "part_of_speech": row.part_of_speech,
                "meaning_zh": row.meaning_zh,
                "is_phrase": row.is_phrase,
                "needs_review": row.needs_review,
                "note": row.note,
                "sources": row.sources,
            }
            for row in by_level[level]
        ]
        (OUTPUT_DIR / f"{level.lower()}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8"
        )
    (OUTPUT_DIR / "cross-level-conflicts.json").write_text(
        json.dumps(conflicts, ensure_ascii=False, indent=1), encoding="utf-8"
    )

    total = sum(len(rows) for rows in by_level.values())
    print(f"提取完成：共 {total} 条（去重后），输出在 {OUTPUT_DIR}")
    for level in LEVEL_ORDER:
        print(f"  {level}: {len(by_level[level])}")
    print(f"跨级冲突: {len(conflicts)}")

    if args.report:
        REPORT_PATH.write_text(build_report(by_level, conflicts), encoding="utf-8")
        print(f"质量报告: {REPORT_PATH}")


if __name__ == "__main__":
    main()
