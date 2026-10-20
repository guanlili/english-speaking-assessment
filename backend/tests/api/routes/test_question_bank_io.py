"""题库批量导入导出：模板/预览/确认/导出全链测试。

覆盖（本期验收口径）：
- 模板下载（xlsx 分工作表可被自家解析器读回、CSV 题型列）；
- 预览行号定位：错误定位到 工作表+行号+原因，原因复用建题校验口径；
- 重复两级：库内/批内完全重复默认跳过；疑似重复确认开关控制是否导入；
- 兜底：问答题主题自动新建 Scenario、单元降级、挂靠篇目三分支；
- 确认导入单事务：注入提交失败整体回滚；
- 导出→导入往返：题干/题型/级别/秒数/话题卡/准备时间等关键配置不丢；
- 权限：学生 403、未登录 401。
"""

import csv
import io
import uuid
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app import crud
from app.models import (
    Instruction,
    Passage,
    RepeatSentence,
    Scenario,
    ScenarioQuestion,
    Unit,
    UserCreate,
)
from app.services import content_dedup
from tests.utils.credential import anonymous, create_student_user, login_headers
from tests.utils.utils import random_email, random_lower_string

QB = "/api/v1/admin/question-bank"

# 基线：本模块首个用例前库内既有的内容行 ID（见 _isolate_content_tables）
_BASELINE_IDS: dict[str, set] | None = None

# CSV 并集列顺序（与模板一致）
CSV_FIELDS = [
    "标题",
    "单元",
    "主题",
    "难度",
    "内容",
    "中文提示",
    "挂靠篇目",
    "秒数",
    "可听次数",
    "考试题型",
    "考试级别",
    "话题卡要点",
    "准备秒数",
    "音频链接",
    "状态",
]
FIELD_TO_HEADER = {
    "title": "标题",
    "unit_title": "单元",
    "topic": "主题",
    "cefr_band": "难度",
    "text": "内容",
    "translation": "中文提示",
    "passage_title": "挂靠篇目",
    "suggested_seconds": "秒数",
    "replay_limit": "可听次数",
    "exam_kind": "考试题型",
    "exam_level": "考试级别",
    "cue_card_bullets": "话题卡要点",
    "prep_seconds": "准备秒数",
    "audio_url": "音频链接",
    "is_active": "状态",
}


BASELINE_TABLES = (
    ScenarioQuestion,
    RepeatSentence,
    Instruction,
    Passage,
    Scenario,
    Unit,
)


@pytest.fixture(autouse=True)
def _isolate_content_tables(db: Session):
    """只清本模块新增的题库行（基线保护式清理）。

    db 是 session 级共享库：其他模块（如 recovery 的 DEMO01 课堂）依赖
    早前种下的内容，绝不能整表清空（共享库污染规约）。首个用例前记下
    既有行 ID，每个用例后只删本模块新增的行。
    """
    global _BASELINE_IDS
    if _BASELINE_IDS is None:
        _BASELINE_IDS = {
            model.__name__: set(db.exec(select(model.id)).all())  # type: ignore[attr-defined]
            for model in BASELINE_TABLES
        }

    yield

    for model in BASELINE_TABLES:
        baseline = _BASELINE_IDS[model.__name__]
        for row in db.exec(select(model)).all():
            if row.id not in baseline:
                db.delete(row)
    db.commit()


def _new_rows(db: Session, model: type) -> list:
    """查询结果里只留本模块新增的行（排除 init_db 种子与既有内容）。"""
    assert _BASELINE_IDS is not None
    baseline = _BASELINE_IDS[model.__name__]
    # getattr+noqa：直接 row.id 会被 ty 拦（object 无该属性），getattr 常量
    # 又被 ruff B009 拦——两者取 getattr + noqa 是唯一交集
    return [
        row
        for row in db.exec(select(model)).all()
        if getattr(row, "id") not in baseline  # noqa: B009
    ]


def _one_new(db: Session, model: type):
    rows = _new_rows(db, model)
    assert len(rows) == 1, f"expected exactly 1 new {model.__name__}, got {len(rows)}"
    return rows[0]


def _login_teacher(db: Session, client: TestClient) -> dict[str, str]:
    email, password = random_email(), random_lower_string()
    crud.create_user(session=db, user_create=UserCreate(email=email, password=password))
    resp = client.post(
        "/api/v1/login/access-token",
        data={"username": email, "password": password},
    )
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def _csv_file(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["题型", *CSV_FIELDS])
    for row in rows:
        qtype = row.get("__type__", "")
        writer.writerow([qtype, *[str(row.get(f, "")) for f in FIELD_TO_HEADER]])
    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes(sheets: dict[str, list[list[Any]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(
    client: TestClient,
    headers: dict[str, str],
    data: bytes,
    *,
    confirm: bool = False,
    include_lines: str = "",
    filename: str = "bank.csv",
    media_type: str = "text/csv",
):
    files = {"file": (filename, data, media_type)}
    if confirm:
        return client.post(
            f"{QB}/import-confirm",
            files=files,
            data={"include_lines": include_lines},
            headers=headers,
        )
    return client.post(f"{QB}/import-preview", files=files, headers=headers)


def _preview(
    client: TestClient, headers: dict[str, str], data: bytes, **kwargs: Any
) -> dict:
    resp = _upload(client, headers, data, **kwargs)
    assert resp.status_code == 200, resp.text
    return resp.json()


PASSAGE_ROW = {
    "__type__": "文章朗读",
    "title": "Morning Routine",
    "unit_title": "",
    "topic": "Daily Life",
    "cefr_band": "B1",
    "text": "I usually wake up at six o'clock and walk to school.",
    "translation": "我通常六点起床走路去上学。",
    "suggested_seconds": "45",
    "is_active": "启用",
}
QUESTION_ROW = {
    "__type__": "情景问答",
    "topic": "Travel",
    "text": "Do you prefer travelling alone or with friends?",
    "translation": "你喜欢独自旅行还是和朋友一起？",
    "suggested_seconds": "20",
}


# ── 模板下载 ─────────────────────────────────────────────────────────


def test_template_download(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    resp = client.get(f"{QB}/import-template?format=xlsx", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.content[:2] == b"PK"
    assert "question-bank-import-template.xlsx" in resp.headers["content-disposition"]
    assert "spreadsheetml" in resp.headers["content-type"]
    # 模板能被自家解析器读回：示例行落在正确的工作表
    from app.services import question_bank_io as svc

    parsed = svc.parse_upload("t.xlsx", resp.content)
    assert {(r.qtype) for r in parsed.rows} == {
        "passage",
        "repeat",
        "question",
        "instruction",
    }

    resp = client.get(f"{QB}/import-template?format=csv", headers=headers)
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("题型,标题,单元,主题,难度,内容")


# ── 预览：计数 / 行号 / 原因 ────────────────────────────────────────


def test_preview_csv_counts_and_line_numbers(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            PASSAGE_ROW,
            {**QUESTION_ROW, "text": "Question A"},
            # 行 4：题目为空
            {"__type__": "情景问答", "topic": "Travel", "text": ""},
            # 行 5：普通题秒数越界（>60）
            {**QUESTION_ROW, "text": "Question B", "suggested_seconds": "61"},
            # 行 6：难度无效
            {"__type__": "文章朗读", "title": "X", "cefr_band": "C1", "text": "hello"},
            # 行 7：题型列无效
            {"__type__": "阅读理解", "text": "zzz"},
        ]
    )
    body = _preview(client, headers, data)
    counts = {s["type"]: s for s in body["sheets"]}
    assert counts["passage"]["valid_count"] == 1
    assert counts["question"]["valid_count"] == 1
    issues = {(i["type"], i["line"]): i["reason"] for i in body["issues"]}
    assert issues[("question", 4)] == "题目不能为空"
    assert issues[("question", 5)] == "建议秒数需在 10–60 之间（考试题型支持长回答）"
    assert issues[("passage", 6)] == "难度无效，可选：A2/B1/B2"
    invalid_type = [i for i in body["issues"] if i["type"] == "-"]
    assert invalid_type[0]["line"] == 7
    assert invalid_type[0]["reason"].startswith("题型列的值无效")
    assert body["total_valid"] == 2


def test_preview_exam_field_validation_reuse(client: TestClient, db: Session) -> None:
    """考试字段校验复用建题口径：有题型必配级别、话题卡仅 Part2、准备时间区间。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            {**QUESTION_ROW, "text": "Q1", "exam_kind": "interview"},  # 缺级别
            {
                **QUESTION_ROW,
                "text": "Q2",
                "exam_kind": "ielts_p1",
                "exam_level": "PET",
                "cue_card_bullets": "what\nwhen",
            },  # 话题卡仅 p2
            {
                **QUESTION_ROW,
                "text": "Q3",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "prep_seconds": "200",
            },  # 准备时间越界
            {
                "__type__": "听句复述",
                "text": "Repeat me.",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
            },  # 复述句不支持 ielts_p2
            {**QUESTION_ROW, "text": "Q4", "exam_level": "PET"},  # 无题型不能配级别
        ]
    )
    body = _preview(client, headers, data)
    reasons = {i["reason"] for i in body["issues"]}
    assert "选择考试题型时需同时标注级别" in reasons
    assert "话题卡要点仅支持 IELTS Part 2 题型" in reasons
    assert "准备时间须在 10–180 秒之间" in reasons
    assert "考试题型无效：听句复述仅支持 toefl_lnr" in reasons
    assert "未选择考试题型时不能单独设置级别/话题卡/准备时间" in reasons
    assert body["total_valid"] == 0


def test_preview_xlsx_sheets(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    data = _xlsx_bytes(
        {
            "文章朗读": [
                [
                    "标题",
                    "单元",
                    "主题",
                    "难度",
                    "正文",
                    "中文提示",
                    "建议秒数",
                    "音频链接",
                    "状态",
                ],
                ["A Trip", "", "Travel", "B1", "We went to the sea.", "", "45", "", ""],
            ],
            "情景问答": [
                [
                    "主题",
                    "英文题目",
                    "中文提示",
                    "作答秒数",
                    "考试题型",
                    "考试级别",
                    "话题卡要点",
                    "准备秒数",
                    "音频链接",
                ],
                ["Travel", "Where did you go?", "", "20", "", "", "", "", ""],
            ],
        }
    )
    body = _preview(client, headers, data, filename="b.xlsx")
    counts = {s["type"]: s for s in body["sheets"]}
    assert counts["passage"]["valid_count"] == 1
    assert counts["question"]["valid_count"] == 1

    unknown = _xlsx_bytes({"随便": [["标题", "正文"], ["a", "b"]]})
    resp = _upload(client, headers, unknown, filename="b.xlsx")
    assert resp.status_code == 422
    assert "无法识别题型" in resp.json()["detail"]

    missing = _xlsx_bytes({"文章朗读": [["标题", "难度"], ["Only title", "B1"]]})
    resp = _upload(client, headers, missing, filename="b.xlsx")
    assert resp.status_code == 422
    assert "缺少必填列：正文" in resp.json()["detail"]


def test_preview_csv_file_level_errors(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    no_type_col = "标题,内容\nT,hello".encode()
    resp = _upload(client, headers, no_type_col)
    assert resp.status_code == 422
    assert "CSV 缺少题型列" in resp.json()["detail"]

    # 行数上限
    rows = [{**QUESTION_ROW, "text": f"Q{i}"} for i in range(2001)]
    resp = _upload(client, headers, _csv_file(rows))
    assert resp.status_code == 422
    assert "行数超过上限" in resp.json()["detail"]

    # GBK 编码兜底（Excel 中文环境「另存为 CSV」默认 GBK）
    gbk = (
        _csv_file([{**QUESTION_ROW, "text": "Encoding check?"}])
        .decode("utf-8")
        .encode("gb18030")
    )
    body = _preview(client, headers, gbk)
    assert body["total_valid"] == 1


def test_audio_url_validation(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Q audio ok",
                "audio_url": "/api/v1/audio/content/a.mp3",
            },
            {**QUESTION_ROW, "text": "Q audio bad", "audio_url": "https://x.com/a.mp3"},
        ]
    )
    body = _preview(client, headers, data)
    reasons = [i["reason"] for i in body["issues"]]
    assert any(r.startswith("音频链接格式不正确") for r in reasons)
    assert body["total_valid"] == 1


# ── 确认导入：落库 / 兜底 / 事务 ────────────────────────────────────


def test_confirm_creates_all_types_and_auto_scenario(
    client: TestClient, db: Session
) -> None:
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            PASSAGE_ROW,
            {
                "__type__": "听句复述",
                "text": "The sea was calm.",
                "translation": "大海很平静。",
                "suggested_seconds": "8",
                "replay_limit": "5",
                "exam_kind": "toefl_lnr",
                "exam_level": "PET",
            },
            {
                **QUESTION_ROW,
                "text": "Describe your favourite place.",
                "suggested_seconds": "120",
                "exam_kind": "IELTS · Part 2",
                "exam_level": "四级",
                "cue_card_bullets": "where it is\nhow often you go",
                "prep_seconds": "60",
            },
            {
                "__type__": "题目说明",
                "title": "Part B",
                "text": "Now talk about the topic.",
                "suggested_seconds": "30",
            },
        ]
    )
    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    sheets = {s["type"]: s for s in body["sheets"]}
    assert sheets["passage"]["created"] == 1
    assert sheets["repeat"]["created"] == 1
    assert sheets["question"]["created"] == 1
    assert sheets["instruction"]["created"] == 1
    assert body["created_topics"] == ["Travel"]

    question = _one_new(db, ScenarioQuestion)
    assert question.exam_kind == "ielts_p2"
    assert question.exam_level == "CET4"
    assert question.cue_card_bullets == ["where it is", "how often you go"]
    assert question.prep_seconds == 60
    assert question.suggested_seconds == 120
    assert question.scenario_id == _one_new(db, Scenario).id

    sentence = _one_new(db, RepeatSentence)
    assert sentence.replay_limit == 5
    assert sentence.exam_kind == "toefl_lnr"
    assert sentence.exam_level == "PET"

    passage = _one_new(db, Passage)
    assert passage.topic == "Daily Life"
    assert passage.slug  # 自动生成
    instruction = _one_new(db, Instruction)
    assert instruction.title == "Part B"

    # 缺音频提示（信息性）
    assert sheets["passage"]["missing_audio_count"] == 1
    assert body["missing_audio_total"] >= 3


def test_confirm_fallbacks_unit_and_passage_linkage(
    client: TestClient, db: Session
) -> None:
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            # 单元不存在 → 降级无单元
            {**PASSAGE_ROW, "unit_title": "不存在的单元"},
            # 挂靠本次导入的篇目 → 挂靠成功
            {
                "__type__": "听句复述",
                "text": "Attached to new passage.",
                "passage_title": "Morning Routine",
            },
            # 挂靠不存在的篇目 → 独立句
            {
                "__type__": "听句复述",
                "text": "No such passage.",
                "passage_title": "Ghost",
            },
        ]
    )
    body = _preview(client, headers, data)
    assert body["unit_fallback_count"] == 1
    assert body["standalone_sentence_count"] == 1

    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text
    passage = _one_new(db, Passage)
    assert passage.unit_id is None
    attached, standalone = _new_rows(db, RepeatSentence)
    assert attached.passage_id == passage.id
    assert standalone.passage_id is None

    # 再建一个同名篇目后，挂靠同名篇目 → 歧义报错
    resp = client.post(
        "/api/v1/admin/passages",
        json={"title": "Morning Routine", "text": "duplicate title body"},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    again = _csv_file(
        [
            {
                "__type__": "听句复述",
                "text": "Ambiguous link.",
                "passage_title": "Morning Routine",
            }
        ]
    )
    body = _preview(client, headers, again)
    assert any("挂靠篇目匹配到多个同名篇目" in i["reason"] for i in body["issues"])


def test_confirm_unit_match_attaches_existing(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    unit = Unit(title="Unit 1", topic="School Life")
    db.add(unit)
    db.commit()
    db.refresh(unit)
    data = _csv_file([{**PASSAGE_ROW, "unit_title": "unit 1"}])  # 大小写不敏感匹配
    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text
    passage = _one_new(db, Passage)
    assert passage.unit_id == unit.id
    assert passage.topic == "School Life"  # topic 派生自单元


def test_confirm_rollback_on_failure(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.services import question_bank_io as svc

    headers = _login_teacher(db, client)
    before_counts = (
        len(_new_rows(db, Passage)),
        len(_new_rows(db, Scenario)),
    )

    def _boom(session: Session) -> None:
        raise RuntimeError("injected failure")

    monkeypatch.setattr(svc, "_commit", _boom)
    resp = _upload(client, headers, _csv_file([QUESTION_ROW]), confirm=True)
    assert resp.status_code == 500
    assert len(_new_rows(db, Passage)) == before_counts[0]
    assert len(_new_rows(db, Scenario)) == before_counts[1]
    assert _new_rows(db, ScenarioQuestion) == []


# ── 重复处理 ─────────────────────────────────────────────────────────


def test_duplicate_tiers_and_reupload_safety(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    file1 = _csv_file([{**QUESTION_ROW, "text": "Same question."}])
    resp = _upload(client, headers, file1, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1

    # 重传同一文件：预览全判重，确认不产生新题
    body = _preview(client, headers, file1)
    assert body["total_valid"] == 0
    counts = body["sheets"][0]
    assert counts["exact_duplicate_count"] == 1
    assert counts["valid_count"] == 0
    resp = _upload(client, headers, file1, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 0
    assert len(_new_rows(db, ScenarioQuestion)) == 1

    # 同题干不同配置 → 疑似重复：默认排除，勾选后导入
    file2 = _csv_file(
        [{**QUESTION_ROW, "text": "Same question.", "suggested_seconds": "30"}]
    )
    body = _preview(client, headers, file2)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    assert body["suspected"][0]["line"] == 2
    resp = _upload(client, headers, file2, confirm=True)
    assert resp.json()["sheets"][0]["skipped_suspected"] == 1
    assert len(_new_rows(db, ScenarioQuestion)) == 1
    resp = _upload(client, headers, file2, confirm=True, include_lines="question:2")
    assert resp.json()["sheets"][0]["created"] == 1
    assert len(_new_rows(db, ScenarioQuestion)) == 2

    # 批内重复：同文件两行一样 → 1 valid + 1 exact dup
    file3 = _csv_file(
        [
            {**QUESTION_ROW, "text": "In-file dup."},
            {**QUESTION_ROW, "text": "In-file dup."},
        ]
    )
    body = _preview(client, headers, file3)
    counts = body["sheets"][0]
    assert counts["valid_count"] == 1
    assert counts["exact_duplicate_count"] == 1

    # 跨主题同题干 → 疑似重复（其他主题已存在）
    file4 = _csv_file([{**QUESTION_ROW, "text": "Same question.", "topic": "Music"}])
    body = _preview(client, headers, file4)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    assert "其他主题" in body["suspected"][0]["reason"]


# ── 导出 → 导入往返 ──────────────────────────────────────────────────


def _delete_new_rows(db: Session) -> None:
    """删除本模块新增的内容行（保留单元：回导还要按单元名匹配）。"""
    for model in (ScenarioQuestion, RepeatSentence, Passage, Instruction, Scenario):
        for row in _new_rows(db, model):
            db.delete(row)
    db.commit()


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export_roundtrip_field_level(
    client: TestClient, db: Session, fmt: str
) -> None:
    """导出 → 删除本批数据 → 回导 → 逐字段比较（评审要求：不能用查重成功
    代替无损验证）。CSV 与 XLSX 双格式各跑一遍。"""
    headers = _login_teacher(db, client)
    db.add(Unit(title="RT Unit", topic="Roundtrip Topic"))
    db.commit()
    source = _csv_file(
        [
            {
                "__type__": "文章朗读",
                "title": "RT Passage",
                "unit_title": "RT Unit",
                "topic": "",  # 挂单元：落库 topic 派生为单元主题
                "cefr_band": "B2",
                "text": "Roundtrip passage body.",
                "translation": "往返篇目。",
                "suggested_seconds": "60",
                "audio_url": "/api/v1/audio/content/rt-p.mp3",
                "is_active": "停用",
            },
            {
                "__type__": "听句复述",
                "text": "RT sentence attached.",
                "translation": "挂靠句。",
                "passage_title": "RT Passage",
                "suggested_seconds": "12",
                "replay_limit": "4",
                "exam_kind": "toefl_lnr",
                "exam_level": "KET",
                "audio_url": "/api/v1/audio/content/rt-s.mp3",
            },
            {
                "__type__": "听句复述",
                "text": "RT sentence standalone.",
            },
            {
                **QUESTION_ROW,
                "topic": "Roundtrip QA",
                "text": "RT question?",
                "translation": "往返题目。",
                "suggested_seconds": "90",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                # 要点含 ; 和 |：拆分只认换行才可逆（评审 P2）
                "cue_card_bullets": "where; why\nwho| what",
                "prep_seconds": "45",
                "audio_url": "/api/v1/audio/content/rt-q.mp3",
            },
            {
                "__type__": "题目说明",
                "title": "RT Part",
                "text": "RT instruction body.",
                "suggested_seconds": "40",
            },
        ]
    )
    resp = _upload(client, headers, source, confirm=True)
    assert resp.status_code == 200, resp.text

    # 先验「重上传导出文件全部判重」（含种子内容，无疑似/错误）
    export = client.get(f"{QB}/export?format={fmt}", headers=headers)
    assert export.status_code == 200, export.text
    body = _preview(
        client,
        headers,
        export.content,
        filename=f"export.{fmt}",
        media_type="text/csv"
        if fmt == "csv"
        else ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    )
    assert body["total_valid"] == 0, body
    # 共享库里早前模块可能留下同名单元：其篇目回判「单元歧义」属预期报错，
    # 不算往损失真——真正的回归信号是「无疑似、零新建」；本批数据的无损
    # 由下一段（删除→回导→逐字段比较）强验证
    assert body["suspected"] == []

    # 删除本批数据后回导：所有行应重建，逐字段无损
    _delete_new_rows(db)
    resp = _upload(
        client,
        headers,
        export.content,
        confirm=True,
        filename=f"export.{fmt}",
        media_type="text/csv"
        if fmt == "csv"
        else ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    )
    assert resp.status_code == 200, resp.text
    created = {s["type"]: s["created"] for s in resp.json()["sheets"]}
    assert created == {"passage": 1, "repeat": 2, "question": 1, "instruction": 1}, (
        resp.json()["sheets"]
    )

    passage = _one_new(db, Passage)
    assert passage.title == "RT Passage"
    assert passage.topic == "Roundtrip Topic"  # 单元派生主题回导不丢
    assert passage.cefr_band == "B2"
    assert passage.text == "Roundtrip passage body."
    assert passage.translation == "往返篇目。"
    assert passage.suggested_seconds == 60
    assert passage.audio_url == "/api/v1/audio/content/rt-p.mp3"
    assert passage.is_active is False
    unit = db.exec(select(Unit).where(Unit.title == "RT Unit")).one()
    assert passage.unit_id == unit.id

    attached = db.exec(
        select(RepeatSentence).where(RepeatSentence.text == "RT sentence attached.")
    ).one()
    standalone = db.exec(
        select(RepeatSentence).where(RepeatSentence.text == "RT sentence standalone.")
    ).one()
    assert attached.passage_id == passage.id  # 挂靠关系回导不丢
    assert attached.translation == "挂靠句。"
    assert attached.suggested_seconds == 12
    assert attached.replay_limit == 4
    assert attached.exam_kind == "toefl_lnr"
    assert attached.exam_level == "KET"
    assert attached.audio_url == "/api/v1/audio/content/rt-s.mp3"
    assert standalone.passage_id is None
    assert standalone.suggested_seconds == 8  # 缺省值
    assert standalone.replay_limit == 3

    question = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "RT question?")
    ).one()
    scenario = db.exec(select(Scenario).where(Scenario.topic == "Roundtrip QA")).one()
    assert question.scenario_id == scenario.id
    assert question.translation == "往返题目。"
    assert question.suggested_seconds == 90
    assert question.exam_kind == "ielts_p2"
    assert question.exam_level == "PET"
    assert question.cue_card_bullets == ["where; why", "who| what"]  # 结构不变
    assert question.prep_seconds == 45
    assert question.audio_url == "/api/v1/audio/content/rt-q.mp3"

    instruction = _one_new(db, Instruction)
    assert instruction.title == "RT Part"
    assert instruction.text == "RT instruction body."
    assert instruction.suggested_seconds == 40


def test_export_xlsx_rejects_overlong_text(client: TestClient, db: Session) -> None:
    """Excel 单元格 32767 上限：超长正文必须显式拒绝并指明 CSV 替代。"""
    headers = _login_teacher(db, client)
    resp = client.post(
        "/api/v1/admin/passages",
        json={"title": "Overlong", "text": "a" * 40000},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    resp = client.get(f"{QB}/export?format=xlsx", headers=headers)
    assert resp.status_code == 422
    assert "32767" in resp.json()["detail"]
    assert "CSV" in resp.json()["detail"]
    # 字段名必须是真实列名而非误引用（曾打印 <function field ...>）
    assert "的text超过" in resp.json()["detail"]
    assert "<function" not in resp.json()["detail"]

    resp = client.get(f"{QB}/export?format=csv", headers=headers)
    assert resp.status_code == 200, resp.text


def test_export_csv_with_filters(client: TestClient, db: Session) -> None:
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            {**QUESTION_ROW, "text": "Travelling question.", "topic": "Travel"},
            {**QUESTION_ROW, "text": "Music question.", "topic": "Music"},
        ]
    )
    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text

    resp = client.get(
        f"{QB}/export?format=csv&types=question&topic=Travel", headers=headers
    )
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("题型,标题,单元,主题,难度,内容")
    assert "Travelling question." in text
    assert "Music question." not in text

    resp = client.get(f"{QB}/export?band=C1", headers=headers)
    assert resp.status_code == 422
    assert "难度无效" in resp.json()["detail"]


def test_unit_passage_reupload_is_exact_duplicate(
    client: TestClient, db: Session
) -> None:
    """查重口径必须用落库后的最终值：挂单元（topic 派生）的篇目重传同文件
    应判完全重复，而不是疑似重复（评审 P2）。"""
    headers = _login_teacher(db, client)
    db.add(Unit(title="U1", topic="Derived Topic"))
    db.commit()
    file = _csv_file(
        [
            {
                "__type__": "文章朗读",
                "title": "Derived Passage",
                "unit_title": "U1",
                "topic": "",
                "text": "Topic follows the unit.",
            }
        ]
    )
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1
    body = _preview(client, headers, file)
    counts = body["sheets"][0]
    assert counts["exact_duplicate_count"] == 1
    assert counts["suspected_duplicate_count"] == 0, body
    assert counts["valid_count"] == 0
    # 完全重复行不开放重复创建：对它提交勾选 = 选择过期 → 422 要求重新预览
    resp = _upload(client, headers, file, confirm=True, include_lines="passage:2")
    assert resp.status_code == 422
    assert "预览结果已变化" in resp.json()["detail"]


def test_audio_difference_is_suspected(client: TestClient, db: Session) -> None:
    """音频属于题目配置：换音频的同题干应判疑似重复（可勾选导入），而不是
    完全重复被静默跳过（评审 P2）。"""
    headers = _login_teacher(db, client)
    first = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Audio differs.",
                "audio_url": "/api/v1/audio/content/a-one.mp3",
            }
        ]
    )
    resp = _upload(client, headers, first, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1

    second = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Audio differs.",
                "audio_url": "/api/v1/audio/content/a-two.mp3",
            }
        ]
    )
    body = _preview(client, headers, second)
    counts = body["sheets"][0]
    assert counts["suspected_duplicate_count"] == 1, body
    assert counts["exact_duplicate_count"] == 0
    resp = _upload(client, headers, second, confirm=True, include_lines="question:2")
    assert resp.json()["sheets"][0]["created"] == 1
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Audio differs.")
    ).all()
    assert sorted(r.audio_url or "" for r in rows) == [
        "/api/v1/audio/content/a-one.mp3",
        "/api/v1/audio/content/a-two.mp3",
    ]


def test_bullets_split_on_newline_only(client: TestClient, db: Session) -> None:
    """话题卡要点只按换行拆：含分号/竖线的要点保持原样（序列化可逆）。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Bullets keep separators.",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "cue_card_bullets": "where; why\nwho| what",
            }
        ]
    )
    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text
    question = db.exec(
        select(ScenarioQuestion).where(
            ScenarioQuestion.text == "Bullets keep separators."
        )
    ).one()
    assert question.cue_card_bullets == ["where; why", "who| what"]


def test_excel_resource_caps(client: TestClient, db: Session) -> None:
    """xlsx 解析资源上限在读遍历前判定：远端格式单元格撑大有效区、超列数、
    跨表累计数据行超限都要立即 422（评审 P1）。"""
    headers = _login_teacher(db, client)
    from openpyxl import Workbook

    passage_header = [
        "标题",
        "单元",
        "主题",
        "难度",
        "正文",
        "中文提示",
        "建议秒数",
        "音频链接",
        "状态",
    ]

    # ① 只有格式、只有一行数据的远端单元格（如 J2010 设过格式）→ 有效区超限
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("文章朗读")
    sheet.append(passage_header)
    sheet.cell(row=2010, column=10, value="格式尾巴")
    buffer = io.BytesIO()
    workbook.save(buffer)
    resp = _upload(
        client,
        headers,
        buffer.getvalue(),
        filename="cap.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert resp.status_code == 422
    assert "有效区" in resp.json()["detail"]

    # ② 列数超限
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("文章朗读")
    sheet.append(passage_header + [""] * 90)
    sheet.append(["Cap", "", "", "B1", "Too many columns.", "", "", "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)
    resp = _upload(
        client,
        headers,
        buffer.getvalue(),
        filename="cap.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert resp.status_code == 422
    assert "列" in resp.json()["detail"]

    # ③ 跨表累计数据行超限（单表都在行上限内：1100 + 1100）
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("文章朗读", "情景问答"):
        sheet = workbook.create_sheet(name)
        sheet.append(passage_header if name == "文章朗读" else ["主题", "英文题目"])
        for i in range(1100):
            if name == "文章朗读":
                sheet.append(
                    [f"Cap {i}", "", "", "B1", "Row cap body.", "", "", "", ""]
                )
            else:
                sheet.append(["Cap", f"Row cap question {i}?"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    resp = _upload(
        client,
        headers,
        buffer.getvalue(),
        filename="cap.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert resp.status_code == 422
    assert "行数超过上限" in resp.json()["detail"]


# ── 第二轮评审回归 ───────────────────────────────────────────────────


def test_csv_guard_roundtrip_reversible(client: TestClient, db: Session) -> None:
    """公式防护前缀可逆：API 建的原内容（含 - / = / 原生引号 开头）导出→
    回导不变形、不新增重复。手写文件中 `'`+危险字符 按.Excel 转义惯例
    剥前缀属设计行为，不在此断言范围。"""
    headers = _login_teacher(db, client)
    raw_texts = [
        "- Explain your choice.",
        "= not a formula",
        # 原文自带的引号：导出须双前缀转义（''），导入只剥一层——
        # 否则原文引号被当前缀误删（评审 P2：'-ed' → -ed'）
        "'-ed' marks the past tense.",
    ]
    for text in raw_texts:
        resp = client.post(
            "/api/v1/admin/instructions",
            json={"text": text, "suggested_seconds": 25},
            headers=headers,
        )
        assert resp.status_code == 200, resp.text

    export = client.get(f"{QB}/export?format=csv", headers=headers)
    assert export.status_code == 200, export.text
    text = export.content.decode("utf-8-sig")
    assert "'- Explain your choice." in text  # 危险字符：单前缀
    assert "''-ed' marks the past tense." in text  # 原生引号：双前缀转义

    body = _preview(client, headers, export.content, filename="export.csv")
    counts = {s["type"]: s for s in body["sheets"]}["instruction"]
    assert counts["exact_duplicate_count"] >= 3, body  # unguard 还原后判重
    assert counts["valid_count"] == 0
    assert counts["suspected_duplicate_count"] == 0

    instructions = {i.text for i in _new_rows(db, Instruction)}
    assert instructions == set(raw_texts)  # 库内原内容不被前缀污染


def test_repeat_ghost_linkage_reupload_exact_dup(
    client: TestClient, db: Session
) -> None:
    """挂靠篇目找不到而降级独立句后，重传同文件应判完全重复（cfg 用最终
    关联），勾选疑似导入也不会建出第二条相同独立句。"""
    headers = _login_teacher(db, client)
    file = _csv_file(
        [
            {
                "__type__": "听句复述",
                "text": "Ghost linked sentence.",
                "passage_title": "Ghost",
            }
        ]
    )
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1

    body = _preview(client, headers, file)
    counts = body["sheets"][0]
    assert counts["exact_duplicate_count"] == 1, body
    assert counts["suspected_duplicate_count"] == 0

    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 0
    assert len(_new_rows(db, RepeatSentence)) == 1


def test_scenario_exact_match_first_and_case_collision(
    client: TestClient, db: Session
) -> None:
    """库内 Travel/travel 并存：精确名优先；查重按主题身份（scenario_id）
    不按主题名——不同主题下的同题是疑似重复（可勾选导入）而非完全重复；
    歧义行优先报错，不被查重静默吞掉。"""
    headers = _login_teacher(db, client)
    upper = Scenario(topic="Travel")
    lower = Scenario(topic="travel")
    db.add(upper)
    db.add(lower)
    db.commit()
    db.refresh(upper)
    db.refresh(lower)

    # 精确命中：文件写 Travel → 进 Travel（不被规范化字典覆盖到 travel）
    file = _csv_file([{**QUESTION_ROW, "topic": "Travel", "text": "Exact match?"}])
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1
    question = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Exact match?")
    ).one()
    assert question.scenario_id == upper.id

    # 不同主题（travel）下的同题：疑似重复而非完全重复，勾选后可导入
    file = _csv_file([{**QUESTION_ROW, "topic": "travel", "text": "Exact match?"}])
    body = _preview(client, headers, file)
    counts = body["sheets"][0]
    assert counts["suspected_duplicate_count"] == 1, body
    assert counts["exact_duplicate_count"] == 0
    assert counts["valid_count"] == 0  # 默认排除
    resp = _upload(client, headers, file, confirm=True, include_lines="question:2")
    assert resp.json()["sheets"][0]["created"] == 1, resp.text
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Exact match?")
    ).all()
    assert {row.scenario_id for row in rows} == {upper.id, lower.id}

    # 歧义行优先报错：内容即使与库内完全相同也不被查重吞掉
    file = _csv_file([{**QUESTION_ROW, "topic": "TRAVEL", "text": "Exact match?"}])
    body = _preview(client, headers, file)
    assert body["total_valid"] == 0
    assert any("多个大小写不同的同名主题" in i["reason"] for i in body["issues"])
    assert all(
        s["exact_duplicate_count"] == 0
        for s in body["sheets"]
        if s["type"] == "question"
    )

    # 精确命中小写主题
    file = _csv_file([{**QUESTION_ROW, "topic": "travel", "text": "Lower exact?"}])
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1
    question = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Lower exact?")
    ).one()
    assert question.scenario_id == lower.id


def _strip_xlsx_dimension(data: bytes) -> bytes:
    """去掉工作表 XML 的 dimension 声明：模拟无尺寸声明的合法 xlsx
    （max_row/max_column 为 None，声明预检必须被逐行实测兜底）。"""
    import re as re_mod
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(
                ".xml"
            ):
                content = re_mod.sub(rb"<dimension[^>]*/>", b"", content)
            zout.writestr(item, content)
    return out.getvalue()


def test_xlsx_column_cap_without_dimension(client: TestClient, db: Session) -> None:
    """无尺寸声明的 100 列文件绕不过列上限：逐行实测列数兜底。"""
    headers = _login_teacher(db, client)
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("文章朗读")
    sheet.append(["标题", "正文"] + [""] * 98)
    sheet.append(["No dimension", "Column cap via row width."])
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = _strip_xlsx_dimension(buffer.getvalue())

    resp = _upload(
        client,
        headers,
        raw,
        filename="nodim.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert resp.status_code == 422
    assert "列" in resp.json()["detail"]
    assert "100" in resp.json()["detail"]


def test_csv_invalid_rows_count_toward_cap(client: TestClient, db: Session) -> None:
    """无效题型行也计入行数上限：2001 行全无效 → 422，不再绕过。"""
    headers = _login_teacher(db, client)
    rows = [{"__type__": "不存在题型", "text": f"Invalid {i}"} for i in range(2001)]
    resp = _upload(client, headers, _csv_file(rows))
    assert resp.status_code == 422
    assert "行数超过上限" in resp.json()["detail"]


def test_csv_huge_field_parses(client: TestClient, db: Session) -> None:
    """超长字段（> csv 默认 128KB 上限）解析不 500：field_size_limit 已抬高。"""
    headers = _login_teacher(db, client)
    huge = "a" * 200_000
    file = _csv_file([{"__type__": "文章朗读", "title": "Huge field", "text": huge}])
    resp = _upload(client, headers, file)
    assert resp.status_code == 200, resp.text
    assert resp.json()["total_valid"] == 1


def test_export_capped_at_reimportable_rows(client: TestClient, db: Session) -> None:
    """导出上限与导入一致（2000）：超出即 422 提示分批——导出的文件必须
    始终可直接回导。"""
    headers = _login_teacher(db, client)
    scenario = Scenario(topic="Export Cap Topic")
    db.add(scenario)
    db.commit()
    db.refresh(scenario)
    for start in (0, 1000):
        items = [
            {"text": f"Cap question {start + i}?", "suggested_seconds": 20}
            for i in range(1000)
        ]
        resp = client.post(
            f"/api/v1/admin/scenarios/{scenario.id}/questions/batch",
            json={"items": items},
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
    # 批量端点单批上限 1000：第 2001 条走单题接口补足
    resp = client.post(
        f"/api/v1/admin/scenarios/{scenario.id}/questions",
        json={"text": "Cap question 2000?", "suggested_seconds": 20},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text

    resp = client.get(f"{QB}/export?types=question", headers=headers)
    assert resp.status_code == 422
    assert "2000" in resp.json()["detail"]
    assert "分批" in resp.json()["detail"]


def test_in_batch_same_text_passages_suspected(client: TestClient, db: Session) -> None:
    """批内两篇不同标题、相同正文：第二篇应判疑似（不只查库内正文）。"""
    headers = _login_teacher(db, client)
    file = _csv_file(
        [
            {
                "__type__": "文章朗读",
                "title": "First title",
                "text": "Identical body across titles.",
            },
            {
                "__type__": "文章朗读",
                "title": "Second title",
                "text": "Identical body across titles.",
            },
        ]
    )
    body = _preview(client, headers, file)
    counts = body["sheets"][0]
    assert counts["valid_count"] == 1
    assert counts["suspected_duplicate_count"] == 1, body
    assert "相同正文已存在于其他篇目" in body["suspected"][0]["reason"]


def test_repeat_same_title_child_passage_not_exact_dup(
    client: TestClient, db: Session
) -> None:
    """同名根篇与旧拆分子篇是不同记录：挂靠按解析后的 passage 身份查重，
    「挂在根篇的同句」与「挂在子篇的同句」互为疑似，不判完全重复。"""
    headers = _login_teacher(db, client)
    root = Passage(
        title="Shared Name", text="Root passage body.", slug="shared-name-root"
    )
    db.add(root)
    db.commit()
    db.refresh(root)
    # 旧拆分子篇：同名，parent 指向根篇（列表只展示根篇，但库内是独立行）
    child = Passage(
        title="Shared Name",
        text="Root passage body.",
        slug="shared-name-child",
        parent_passage_id=root.id,
    )
    db.add(child)
    db.commit()
    db.refresh(child)
    db.add(
        RepeatSentence(
            passage_id=child.id, order_index=1, text="Linked to child passage."
        )
    )
    db.commit()

    # 挂靠 Shared Name 解析到根篇（唯一根篇命中）——与子篇下的同句身份不同
    file = _csv_file(
        [
            {
                "__type__": "听句复述",
                "text": "Linked to child passage.",
                "passage_title": "Shared Name",
            }
        ]
    )
    body = _preview(client, headers, file)
    counts = body["sheets"][0]
    assert counts["suspected_duplicate_count"] == 1, body
    assert counts["exact_duplicate_count"] == 0
    assert counts["valid_count"] == 0  # 默认排除
    # 勾选后可导入：挂到根篇，与子篇下的那句并存
    resp = _upload(client, headers, file, confirm=True, include_lines="repeat:2")
    assert resp.json()["sheets"][0]["created"] == 1, resp.text
    db.refresh(root)
    root_sentence = db.exec(
        select(RepeatSentence).where(RepeatSentence.text == "Linked to child passage.")
    ).all()
    assert {s.passage_id for s in root_sentence} == {root.id, child.id}


def test_export_size_capped_at_upload_limit(client: TestClient, db: Session) -> None:
    """导出体积对齐导入 5MB 上限：超出即 422 提示分批，不让导出文件回导被 413 挡下。"""
    headers = _login_teacher(db, client)
    # 每篇约 3.3KB 正文 × 1600 行 ≈ 5.3MB > 5MB；行数 1600 < 2000（先触发行数以外的体积门）。
    # 种子直写 DB：导入通道本身受 5MB 上传限制，不可能经导入造出超大数据
    for i in range(1600):
        db.add(
            Passage(
                title=f"Bulk passage {i}",
                text="x" * 3300,
                slug=f"bulk-passage-{i}",
            )
        )
    db.commit()

    # CSV 明文超 5MB → 422 提示分批；xlsx 压缩后实际体积小 → 合法导出，
    # 但同样必须过体积门（断言其物理尺寸 ≤ 5MB，保证可回导）
    resp = client.get(f"{QB}/export?types=passage&format=csv", headers=headers)
    assert resp.status_code == 422, resp.status_code
    assert "5MB" in resp.json()["detail"], resp.json()["detail"]
    assert "分批" in resp.json()["detail"]

    resp = client.get(f"{QB}/export?types=passage&format=xlsx", headers=headers)
    assert resp.status_code == 200, resp.text
    assert len(resp.content) <= 5 * 1024 * 1024

    # 缩小范围后可正常导出（体积门按筛选后的实际内容判定）
    resp = client.get(
        f"{QB}/export?types=passage&q=Bulk%20passage%20159&format=csv",
        headers=headers,
    )
    assert resp.status_code == 200, resp.text


# ── 第四轮：导入后直达使用 / 错误下载 / 逐条确认 / 并发防重 ─────────


def test_confirm_returns_created_items(client: TestClient, db: Session) -> None:
    """确认导入返回本批新建条目（真实 ID/label/topic/音频需求），零新建时为空列表。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            PASSAGE_ROW,
            {"__type__": "听句复述", "text": "Needs audio sentence."},
            {
                **QUESTION_ROW,
                "text": "Created item question?",
                "audio_url": "/api/v1/audio/content/q.mp3",
            },
            {
                "__type__": "题目说明",
                "title": "CI Part",
                "text": "Created item instruction.",
            },
        ]
    )
    resp = _upload(client, headers, data, confirm=True)
    assert resp.status_code == 200, resp.text
    items = resp.json()["created_items"]
    assert len(items) == 4, resp.json()

    by_type = {item["type"]: item for item in items}
    assert by_type["passage"]["label"] == PASSAGE_ROW["title"]
    assert by_type["passage"]["needs_audio"] is False
    assert by_type["repeat"]["needs_audio"] is True
    assert by_type["repeat"]["audio_missing"] is True
    assert by_type["repeat"]["label"].startswith("Needs audio")
    assert by_type["question"]["topic"] == QUESTION_ROW["topic"]
    assert by_type["question"]["audio_missing"] is False
    assert by_type["instruction"]["label"] == "CI Part"
    # id 是库里真实行：按 id 能查回
    sentence = db.get(RepeatSentence, uuid.UUID(by_type["repeat"]["id"]))
    assert sentence is not None and sentence.text == "Needs audio sentence."

    # 零新增（全部重复）：created_items 为空，不退化成全库
    resp = _upload(client, headers, data, confirm=True)
    assert resp.json()["created_items"] == []


def test_error_file_roundtrip_and_full_list(client: TestClient, db: Session) -> None:
    """下载待修正数据：原始字段+行号+原因全保留（含原生引号），修正后直接
    回传成功（辅助列被忽略）；错误清单不受 200 条界面截断限制。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            {**QUESTION_ROW, "text": "Valid row."},
            {"__type__": "不存在题型", "text": "Bad type row."},  # 解析期错误
            {
                "__type__": "情景问答",
                "topic": "FixTopic",
                "text": "x" * 600,  # 题目过长
                "translation": "'-ed' 原生引号保留检查",
            },
            {
                "__type__": "情景问答",
                "topic": "FixTopic",
                "text": "Seconds out of range?",
                "suggested_seconds": "999",  # 秒数越界
            },
        ]
    )
    body = _preview(client, headers, data)
    counts = {sheet["type"]: sheet for sheet in body["sheets"]}["question"]
    assert counts["invalid_count"] == 2
    assert body["issues"][0]["type"] == "-"  # 解析期错误行也在清单

    resp = client.post(
        f"{QB}/import-error-file",
        files={"file": ("bank.csv", data, "text/csv")},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    lines = text.splitlines()
    assert lines[0].endswith("原行号,错误原因")  # 辅助列在表头
    assert len(lines) == 4  # 表头 + 3 条错误行（含解析期错误）
    assert "x" * 60 in lines[2]  # 原始字段保留
    # 手写文件中 '+危险字符 按 Excel 转义惯例剥离（设计口径）；错误文件对
    # 剥离后的规范值再加防护前缀——不改字段直接回传，值保持稳定不漂移
    assert "'-ed' 原生引号保留检查" in text
    assert "999" in lines[3]
    assert "题目过长" in lines[2] and "行数超过上限" not in "".join(lines)

    # 模拟老师修正：改对三处错误，辅助列原样保留
    import csv as csv_mod
    import io as io_mod

    rows = list(csv_mod.reader(io_mod.StringIO(text.lstrip("\ufeff"))))
    header, fix_rows = rows[0], rows[1:]
    fix_rows[0][0] = "情景问答"  # 题型改对（辅助列在尾部不影响）
    fix_rows[0][3] = "FixTopic"  # 主题补上
    fix_rows[0][5] = "Fixed type row?"
    fix_rows[1][5] = "Shortened question?"
    fix_rows[2][8] = "20"  # 秒数改对（第 9 列）
    fixed = io_mod.StringIO()
    writer = csv_mod.writer(fixed)
    writer.writerows([header, *fix_rows])
    resp = _upload(client, headers, fixed.getvalue().encode("utf-8"), confirm=True)
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 3, resp.text  # 修正后全部入库
    assert resp.json()["created_items"]
    # 引号原文在库内一字不差
    question = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Shortened question?")
    ).one()
    assert question.translation == "-ed' 原生引号保留检查"

    # 无错误行 → 422 明确提示
    clean = _csv_file([{**QUESTION_ROW, "text": "Clean row."}])
    resp = client.post(
        f"{QB}/import-error-file",
        files={"file": ("bank.csv", clean, "text/csv")},
        headers=headers,
    )
    assert resp.status_code == 422
    assert "没有需要修正的行" in resp.json()["detail"]


def test_error_file_full_list_beyond_200(client: TestClient, db: Session) -> None:
    """201 条错误行：界面 issues 截断到 200，下载的是全量。"""
    headers = _login_teacher(db, client)
    rows = [
        {"__type__": "情景问答", "topic": "Bulk", "text": f"Too long {i} " + "x" * 600}
        for i in range(201)
    ]
    data = _csv_file(rows)
    body = _preview(client, headers, data)
    assert body["issues_truncated"] is True

    resp = client.post(
        f"{QB}/import-error-file",
        files={"file": ("bank.csv", data, "text/csv")},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert len(resp.content.decode("utf-8-sig").splitlines()) == 202  # 表头 + 201


def test_suspected_snapshot_and_per_row_selection(
    client: TestClient, db: Session
) -> None:
    """疑似快照展示两侧关键字段差异；逐条选择只导入勾选行，默认全跳过。"""
    headers = _login_teacher(db, client)
    seed = _csv_file(
        [{**QUESTION_ROW, "text": "Diff question?", "suggested_seconds": "20"}]
    )
    resp = _upload(client, headers, seed, confirm=True)
    assert resp.status_code == 200, resp.text

    # 同题干不同秒数 → 疑似：快照须含两侧对照
    file = _csv_file(
        [
            {**QUESTION_ROW, "text": "Diff question?", "suggested_seconds": "30"},
            {**QUESTION_ROW, "text": "Fresh valid question?"},
        ]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    suspected = body["suspected"][0]
    assert suspected["type"] == "question" and suspected["line"] == 2
    assert suspected["incoming"]["seconds"] == "30"
    assert suspected["existing"]["seconds"] == "20"
    assert suspected["incoming"]["text"] == suspected["existing"]["text"]

    # 默认全跳过：只导入 1 条有效行
    resp = _upload(client, headers, file, confirm=True)
    sheet = resp.json()["sheets"][0]
    assert sheet["created"] == 1
    assert sheet["skipped_suspected"] == 1
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Diff question?")
    ).all()
    assert len(rows) == 1 and rows[0].suggested_seconds == 20

    # 换库状态后（把疑似行变成可再导）：勾选 line 2 精确导入该行
    file2 = _csv_file(
        [{**QUESTION_ROW, "text": "Diff question?", "suggested_seconds": "40"}]
    )
    resp = _upload(client, headers, file2, confirm=True, include_lines="question:2")
    sheet = resp.json()["sheets"][0]
    assert sheet["created"] == 1
    assert sheet["skipped_suspected"] == 0
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Diff question?")
    ).all()
    assert sorted(r.suggested_seconds for r in rows) == [20, 40]


def test_stale_selection_rejected(client: TestClient, db: Session) -> None:
    """预览后库被他人写入 → 疑似行变成完全重复：勾选令牌过期 → 422 重新预览。"""
    headers = _login_teacher(db, client)
    other_headers = _login_teacher(db, client)
    file = _csv_file(
        [{**QUESTION_ROW, "text": "Race question?", "suggested_seconds": "20"}]
    )
    # 老师A 先导入一版（不同配置），使老师B 的文件成为疑似
    seeded = _csv_file(
        [{**QUESTION_ROW, "text": "Race question?", "suggested_seconds": "30"}]
    )
    assert _upload(client, other_headers, seeded, confirm=True).status_code == 200

    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1

    # 预览与确认之间，另一账号勾选导入了**完全相同**配置 → 疑似变完全重复
    same = _csv_file(
        [{**QUESTION_ROW, "text": "Race question?", "suggested_seconds": "20"}]
    )
    resp_other = _upload(
        client, other_headers, same, confirm=True, include_lines="question:2"
    )
    assert resp_other.status_code == 200, resp_other.text

    resp = _upload(client, headers, file, confirm=True, include_lines="question:2")
    assert resp.status_code == 422
    assert "预览结果已变化" in resp.json()["detail"]
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Race question?")
    ).all()
    assert len(rows) == 2  # 20s + 30s，无第三条


def test_concurrent_confirm_two_connections(db: Session, client: TestClient) -> None:
    """两个独立数据库连接并发确认同一批：咨询锁串行化 + 锁内重查，
    只有一方落库，另一方全部判重跳过；主题不重复、无 500。"""
    import threading

    from sqlalchemy import create_engine

    from app.core.config import settings
    from app.services import question_bank_io as svc

    parsed = svc.parse_upload(
        "race.csv",
        _csv_file(
            [
                {**QUESTION_ROW, "text": "Concurrent question?", "topic": "RaceTopic"},
                {**QUESTION_ROW, "text": "Second concurrent?", "topic": "RaceTopic"},
            ]
        ),
    )
    engine = create_engine(str(settings.SQLALCHEMY_DATABASE_TEST_URI))
    results: list[Any] = []
    errors: list[BaseException] = []

    def _run() -> None:
        import sqlmodel

        with sqlmodel.Session(engine) as session:
            try:
                results.append(svc.apply_import(session, parsed))
            except BaseException as exc:  # noqa: BLE001
                errors.append(exc)

    threads = [threading.Thread(target=_run) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors, errors

    created_totals = sorted(
        sum(sheet.created for sheet in result.sheets) for result in results
    )
    assert created_totals == [0, 2], created_totals  # 一方全建、一方全跳过
    scenarios = db.exec(select(Scenario).where(Scenario.topic == "RaceTopic")).all()
    assert len(scenarios) == 1
    from sqlmodel import col

    questions = db.exec(
        select(ScenarioQuestion).where(
            col(ScenarioQuestion.text).in_(  # type: ignore[union-attr]
                ["Concurrent question?", "Second concurrent?"]
            )
        )
    ).all()
    assert len(questions) == 2


def test_integrity_conflict_returns_409(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    """陈旧计划撞唯一约束（锁内重查失效的极端兜底）：409 而非 500，整批回滚。"""
    from app.services import question_bank_io as svc

    headers = _login_teacher(db, client)
    file = _csv_file(
        [{**QUESTION_ROW, "text": "Conflict question?", "topic": "ConflictTopic"}]
    )
    # 造一份「查重为空」的陈旧计划（模拟锁内重查前库内状态）
    stale_plan = svc.plan_import(db, svc.parse_upload("stale.csv", file))
    assert stale_plan.new_topics == ["ConflictTopic"]
    # 计划生成后，库内出现同名主题 → 陈旧计划插入必撞 topic 唯一约束
    db.add(Scenario(topic="ConflictTopic"))
    db.commit()
    monkeypatch.setattr(svc, "plan_import", lambda *args, **kwargs: stale_plan)

    resp = _upload(client, headers, file, confirm=True)
    assert resp.status_code == 409, resp.text
    assert "冲突" in resp.json()["detail"]
    # 整批回滚：没有多余主题、没有题目
    assert (
        len(db.exec(select(Scenario).where(Scenario.topic == "ConflictTopic")).all())
        == 1
    )
    assert (
        db.exec(
            select(ScenarioQuestion).where(
                ScenarioQuestion.text == "Conflict question?"
            )
        ).first()
        is None
    )


# ── 第五轮：指纹对账 / 同题型双表 / 对照字段 / 疑似全量 / 音频规则 ────


def test_fingerprint_rejects_content_change(client: TestClient, db: Session) -> None:
    """预览后候选题内容被改：指纹失配 → 422 重新预览（此前只查「还是不是
    疑似」，候选 20→60 秒照样放行）。"""
    headers = _login_teacher(db, client)
    seeded = _csv_file(
        [{**QUESTION_ROW, "text": "FP question?", "suggested_seconds": "20"}]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    file = _csv_file(
        [{**QUESTION_ROW, "text": "FP question?", "suggested_seconds": "30"}]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    assert body["fingerprint"]
    fingerprint = body["fingerprint"]

    # 预览与确认之间：候选题（20 秒那道）被编辑为 60 秒 → 对照快照变化
    existing = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "FP question?")
    ).one()
    existing.suggested_seconds = 60
    db.add(existing)
    db.commit()

    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={
            "include_lines": "question:2",
            "preview_fingerprint": fingerprint,
        },
        headers=headers,
    )
    assert resp.status_code == 422, resp.text
    assert "预览结果已变化" in resp.json()["detail"]
    assert (
        db.exec(select(ScenarioQuestion).where(ScenarioQuestion.text == "FP question?"))
        .all()
        .__len__()
        == 1
    )  # 未新增

    # 重新预览（新指纹反映 60 秒候选）后同一文件可正常勾选导入
    body = _preview(client, headers, file)
    assert body["suspected"][0]["existing"]["seconds"] == "60"
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={
            "include_lines": "question:2",
            "preview_fingerprint": body["fingerprint"],
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 1


def test_same_qtype_two_sheets_rejected(client: TestClient, db: Session) -> None:
    """「情景问答」和「question」两表同题型：行标识 (type,line) 会撞号 →
    文件级 422，要求合并到一个工作表。"""
    headers = _login_teacher(db, client)
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("情景问答", "question"):
        sheet = workbook.create_sheet(name)
        sheet.append(["主题", "英文题目"])
        sheet.append(["Dup", f"Sheet {name} row?"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    resp = _upload(
        client,
        headers,
        buffer.getvalue(),
        filename="dup-type.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert resp.status_code == 422
    assert "多个工作表" in resp.json()["detail"]


def test_suspected_snapshot_includes_translation_and_cue(
    client: TestClient, db: Session
) -> None:
    """只有中文提示/话题卡不同的两道题：对照快照必须把差异字段亮出来。"""
    headers = _login_teacher(db, client)
    first = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Diff by hint only?",
                "translation": "第一版提示",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "cue_card_bullets": "who\nwhy",
            }
        ]
    )
    assert _upload(client, headers, first, confirm=True).status_code == 200

    second = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Diff by hint only?",
                "translation": "第二版提示",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "cue_card_bullets": "who\nwhere",
            }
        ]
    )
    body = _preview(client, headers, second)
    suspected = body["suspected"][0]
    assert suspected["incoming"]["translation"] == "第二版提示"
    assert suspected["existing"]["translation"] == "第一版提示"
    assert suspected["incoming"]["cue"] == "who / where"
    assert suspected["existing"]["cue"] == "who / why"


def test_suspected_list_full_beyond_100(client: TestClient, db: Session) -> None:
    """101 条疑似全量返回：全部导入可一次勾完（截断会漏掉后续行）。"""
    headers = _login_teacher(db, client)
    rows = [
        {**QUESTION_ROW, "text": f"Bulk seeded {i}?", "suggested_seconds": "20"}
        for i in range(101)
    ]
    assert _upload(client, headers, _csv_file(rows), confirm=True).status_code == 200

    changed = _csv_file(
        [
            {**QUESTION_ROW, "text": f"Bulk seeded {i}?", "suggested_seconds": "30"}
            for i in range(101)
        ]
    )
    body = _preview(client, headers, changed)
    assert body["sheets"][0]["suspected_duplicate_count"] == 101
    assert len(body["suspected"]) == 101, "疑似清单不得截断"
    assert body["suspected_truncated"] is False

    include_lines = ",".join(f"question:{2 + i}" for i in range(101))
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", changed, "text/csv")},
        data={
            "include_lines": include_lines,
            "preview_fingerprint": body["fingerprint"],
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 101


def test_created_items_needs_audio_matches_exam_rule(
    client: TestClient, db: Session
) -> None:
    """needs_audio 对齐自动考试规则：所有非朗读题（复述/问答/说明）为 True。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [
            PASSAGE_ROW,
            {"__type__": "听句复述", "text": "Exam audio rule sentence."},
            {**QUESTION_ROW, "text": "Exam audio rule question?"},
            {"__type__": "题目说明", "text": "Exam audio rule instruction."},
        ]
    )
    resp = _upload(client, headers, data, confirm=True)
    items = {item["type"]: item for item in resp.json()["created_items"]}
    assert items["passage"]["needs_audio"] is False
    assert items["repeat"]["needs_audio"] is True
    assert items["question"]["needs_audio"] is True
    assert items["instruction"]["needs_audio"] is True
    assert items["instruction"]["audio_missing"] is True  # exam_settings 导入不写


def test_preview_invalid_total_covers_parse_errors(
    client: TestClient, db: Session
) -> None:
    """全部为无效题型行：invalid_total>0（下载入口依据），不依赖题型计数。"""
    headers = _login_teacher(db, client)
    data = _csv_file(
        [{"__type__": "不存在题型", "text": f"All bad {i}"} for i in range(3)]
    )
    body = _preview(client, headers, data)
    assert body["invalid_total"] == 3
    assert body["total_valid"] == 0


# ── 第六轮：指纹决策无关 / 尾部敏感 / 空题单防护 ─────────────────────


def test_fingerprint_independent_of_selection(client: TestClient, db: Session) -> None:
    """勾选疑似行不得改变指纹：新主题导入同题（预览疑似 → 勾选导入），
    确认必须成功——此前 new_topics 受决策影响，确认永远 422。"""
    headers = _login_teacher(db, client)
    seeded = _csv_file(
        [{**QUESTION_ROW, "text": "Decision independent?", "topic": "BaseTopic"}]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    file = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Decision independent?",
                "topic": "FreshTopic",
            }
        ]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    # 同一指纹反复预览稳定
    fingerprint = body["fingerprint"]
    body2 = _preview(client, headers, file)
    assert body2["fingerprint"] == fingerprint

    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "question:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 1
    assert resp.json()["created_topics"] == ["FreshTopic"]


def test_fingerprint_detects_trailing_change(client: TestClient, db: Session) -> None:
    """候选内容尾部（60 字符之后）被改：指纹必须失配（此前摘要截断放行）。"""
    headers = _login_teacher(db, client)
    long_hint = "提" * 80
    seeded = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Trailing change?",
                "translation": long_hint,
            }
        ]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    file = _csv_file(
        [{**QUESTION_ROW, "text": "Trailing change?", "translation": "改"}]
    )
    body = _preview(client, headers, file)
    fingerprint = body["fingerprint"]
    # 对照快照携带完整值（含 80 字提示），不是 60 字截断
    assert body["suspected"][0]["existing"]["translation"] == long_hint

    existing = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Trailing change?")
    ).one()
    existing.translation = long_hint[:70]  # 尾部第 70 字改动（旧摘要截断盲区）
    db.add(existing)
    db.commit()

    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "question:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 422
    assert "预览结果已变化" in resp.json()["detail"]


# ── 第七轮：同批双疑似全量导入 / 指纹候选全量 ─────────────────────────


def test_same_batch_two_suspected_import_all(client: TestClient, db: Session) -> None:
    """同批两条相同题干均疑似：「全部导入」不再误判预览过期——决策产生的
    完全重复按 exact 正常跳过（评审 P2 复现：此前整批 422）。"""
    headers = _login_teacher(db, client)
    seeded = _csv_file(
        [{**QUESTION_ROW, "text": "Twin suspected?", "suggested_seconds": "20"}]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    # 变体一：两行配置不同（30/40s）→ 预览两条均疑似，全部导入 → 两条入库
    file = _csv_file(
        [
            {**QUESTION_ROW, "text": "Twin suspected?", "suggested_seconds": "30"},
            {**QUESTION_ROW, "text": "Twin suspected?", "suggested_seconds": "40"},
        ]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 2, body
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={
            "include_lines": "question:2,question:3",
            "preview_fingerprint": body["fingerprint"],
        },
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 2, resp.text

    # 变体二：两行配置完全相同 → 预览第一条 valid、第二条批内 exact；
    # 无疑似可勾，确认（无勾选）只建一条，绝不 422
    file2 = _csv_file(
        [
            {**QUESTION_ROW, "text": "Identical twins?", "suggested_seconds": "50"},
            {**QUESTION_ROW, "text": "Identical twins?", "suggested_seconds": "50"},
        ]
    )
    body2 = _preview(client, headers, file2)
    sheet = body2["sheets"][0]
    assert sheet["valid_count"] == 1, sheet
    assert sheet["exact_duplicate_count"] == 1  # 第二条批内完全重复
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file2, "text/csv")},
        data={"preview_fingerprint": body2["fingerprint"]},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["sheets"][0]["created"] == 1
    rows = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Identical twins?")
    ).all()
    assert len(rows) == 1


def test_fingerprint_covers_all_candidates(client: TestClient, db: Session) -> None:
    """指纹覆盖全部冲突候选：同题第二条候选被新增或修改都要失配
    （此前只取展示代表，第二条变化放行）。"""
    headers = _login_teacher(db, client)
    # 库内同主题两条同题（20s / 25s）
    scenario = Scenario(topic="MultiCandidate")
    db.add(scenario)
    db.commit()
    db.refresh(scenario)
    for seconds in (20, 25):
        db.add(
            ScenarioQuestion(
                scenario_id=scenario.id,
                order_index=seconds,
                text="Multi candidate?",
                suggested_seconds=seconds,
            )
        )
    db.commit()

    file = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Multi candidate?",
                "topic": "MultiCandidate",
                "suggested_seconds": "30",
            }
        ]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    fingerprint = body["fingerprint"]

    # (a) 预览后**新增**第三条同题候选（28s）→ 旧指纹失配
    db.add(
        ScenarioQuestion(
            scenario_id=scenario.id,
            order_index=3,
            text="Multi candidate?",
            suggested_seconds=28,
        )
    )
    db.commit()
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "question:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 422, resp.text
    assert "预览结果已变化" in resp.json()["detail"]

    # 重新预览拿新指纹；改**第二条**候选（25→26s）→ 再失配
    body = _preview(client, headers, file)
    fingerprint = body["fingerprint"]
    second = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.suggested_seconds == 25)
    ).one()
    second.suggested_seconds = 26
    db.add(second)
    db.commit()
    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "question:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 422, resp.text


# ── 第八轮：指纹覆盖启停状态与话题卡结构 ───────────────────────────────


def test_fingerprint_detects_active_toggle(client: TestClient, db: Session) -> None:
    """候选篇目被停用/启用：指纹失配（启停参与查重配置，也须参与指纹）。"""
    headers = _login_teacher(db, client)
    seeded = _csv_file(
        [
            {
                "__type__": "文章朗读",
                "title": "Toggle me",
                "text": "Toggle candidate body.",
                "suggested_seconds": "30",
            }
        ]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    file = _csv_file(
        [
            {
                "__type__": "文章朗读",
                "title": "Toggle me",
                "text": "Toggle candidate body.",
                "suggested_seconds": "45",
            }
        ]
    )
    body = _preview(client, headers, file)
    assert body["sheets"][0]["suspected_duplicate_count"] == 1
    fingerprint = body["fingerprint"]

    existing = db.exec(select(Passage).where(Passage.title == "Toggle me")).one()
    existing.is_active = not existing.is_active
    db.add(existing)
    db.commit()

    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "passage:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 422
    assert "预览结果已变化" in resp.json()["detail"]


def test_fingerprint_detects_bullets_structure_change(
    client: TestClient, db: Session
) -> None:
    """话题卡数组结构变化（两条合并为一条，join 后文本相同）：指纹失配。"""
    headers = _login_teacher(db, client)
    seeded = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Bullets structure?",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "cue_card_bullets": "who\nwhy",
            }
        ]
    )
    assert _upload(client, headers, seeded, confirm=True).status_code == 200

    file = _csv_file(
        [
            {
                **QUESTION_ROW,
                "text": "Bullets structure?",
                "exam_kind": "ielts_p2",
                "exam_level": "PET",
                "cue_card_bullets": "who\nwhy\nwhere",
            }
        ]
    )
    body = _preview(client, headers, file)
    fingerprint = body["fingerprint"]

    existing = db.exec(
        select(ScenarioQuestion).where(ScenarioQuestion.text == "Bullets structure?")
    ).one()
    # join 展示相同（"who / why"），数组结构不同：["who","why"] → ["who / why"]
    existing.cue_card_bullets = ["who / why"]
    db.add(existing)
    db.commit()

    resp = client.post(
        f"{QB}/import-confirm",
        files={"file": ("bank.csv", file, "text/csv")},
        data={"include_lines": "question:2", "preview_fingerprint": fingerprint},
        headers=headers,
    )
    assert resp.status_code == 422
    assert "预览结果已变化" in resp.json()["detail"]


# ── 权限 ─────────────────────────────────────────────────────────────


def test_permissions_student_403_and_anonymous_401(
    client: TestClient, db: Session
) -> None:
    student = create_student_user(db)
    assert student.username is not None
    student_headers = login_headers(client, student.username)
    resp = client.get(f"{QB}/import-template", headers=student_headers)
    assert resp.status_code == 403
    with anonymous(client):
        resp = client.get(f"{QB}/import-template")
        assert resp.status_code == 401


def test_repeat_import_streams_unicode_keys_without_materializing_unrelated_tables(
    db: Session,
) -> None:
    from sqlalchemy import event

    from app.services import question_bank_io as svc

    marker = uuid.uuid4().hex
    passage = Passage(
        slug=f"candidate-{marker}",
        title="Unrelated",
        text=f"Unrelated {marker}",
        cefr_band="B1",
    )
    instruction = Instruction(text=f"Unrelated {marker}", suggested_seconds=20)
    sentence = RepeatSentence(
        text=f"Straße\u00a0ＦＵＬＬ {marker}",
        order_index=0,
        suggested_seconds=8,
        replay_limit=3,
    )
    db.add_all([passage, instruction, sentence])
    db.commit()
    ids = passage.id, instruction.id, sentence.id
    for row in (passage, instruction, sentence):
        db.expunge(row)
    loaded: list[type] = []

    def remember(_session: Session, instance: object) -> None:
        loaded.append(type(instance))

    event.listen(db, "loaded_as_persistent", remember)
    try:
        file = _csv_file(
            [
                {
                    "__type__": "听句复述",
                    "text": f"STRASSE FULL {marker}",
                    "passage_title": "",
                }
            ]
        )
        plan = svc.plan_import(db, svc.parse_upload("unicode.csv", file))
        assert plan.counts[svc.TYPE_REPEAT].exact_dup == 1
        assert RepeatSentence in loaded
        assert Passage not in loaded
        assert Instruction not in loaded
        assert ScenarioQuestion not in loaded
    finally:
        event.remove(db, "loaded_as_persistent", remember)
        for model, row_id in zip(
            (Passage, Instruction, RepeatSentence), ids, strict=True
        ):
            row = db.get(model, row_id)
            if row is not None:
                db.delete(row)
        db.commit()


def test_dedup_hashes_maintained_on_write_and_update(
    client: TestClient, db: Session
) -> None:
    """flush 钩子维护 dedup_*_hash：创建即写入、改文本随更新重算，
    归一化口径（全角空白/大小写折叠）与 content_dedup.norm_text 一致。"""
    headers = _login_teacher(db, client)
    file = _csv_file([{**QUESTION_ROW, "text": "Hash maintained."}])
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1
    row = _one_new(db, ScenarioQuestion)
    scenario_id = row.scenario_id
    try:
        assert row.dedup_text_hash == content_dedup.hash_text("Hash maintained.")

        row.text = "Hash maintained v2."
        db.add(row)
        db.commit()
        db.refresh(row)
        assert row.dedup_text_hash == content_dedup.hash_text("Hash maintained v2.")

        # 全角空格 + 大写 + 多重空白 → 归一后与首版同键
        row.text = "　Hash　MAINTAINED.  "
        db.add(row)
        db.commit()
        db.refresh(row)
        assert row.dedup_text_hash == content_dedup.hash_text("Hash maintained.")
    finally:
        db.delete(row)
        db.commit()
        scenario = db.get(Scenario, scenario_id)
        if scenario is not None:
            db.delete(scenario)
            db.commit()


def test_matching_falls_back_to_stream_when_hash_missing(
    client: TestClient, db: Session
) -> None:
    """哈希列为 NULL 的存量行（迁移前残留/绕过 ORM 的写入）仍走流式归一
    判重，语义与预计算路径一致。"""
    from sqlalchemy import text

    headers = _login_teacher(db, client)
    file = _csv_file([{**QUESTION_ROW, "text": "Legacy unhashed."}])
    resp = _upload(client, headers, file, confirm=True)
    assert resp.json()["sheets"][0]["created"] == 1
    row = _one_new(db, ScenarioQuestion)
    scenario_id = row.scenario_id
    try:
        db.expunge(row)
        # connection() 走 SQLAlchemy Connection 层执行原生 SQL：
        # session.execute 被 ty 按弃用拦截
        conn = db.connection()
        conn.execute(
            text("UPDATE scenario_question SET dedup_text_hash = NULL WHERE id = :id"),
            {"id": row.id},
        )
        db.commit()

        body = _preview(client, headers, file)
        counts = body["sheets"][0]
        assert counts["exact_duplicate_count"] == 1
        assert counts["valid_count"] == 0
    finally:
        db.connection().execute(
            text("DELETE FROM scenario_question WHERE id = :id"), {"id": row.id}
        )
        db.commit()
        scenario = db.get(Scenario, scenario_id)
        if scenario is not None:
            db.delete(scenario)
            db.commit()
